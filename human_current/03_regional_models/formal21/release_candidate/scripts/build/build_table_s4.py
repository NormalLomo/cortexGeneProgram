#!/usr/bin/env python3
"""Build the donor-aware regional-program Supplementary Table S4 from fixed inputs."""

import os
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


ROOT_VALUE = os.environ.get("FORMAL21_PROJECT_ROOT")
if not ROOT_VALUE:
    raise RuntimeError("FORMAL21_PROJECT_ROOT is required")
ROOT = Path(ROOT_VALUE).resolve(strict=True)
RESULTS = ROOT / "results" / "crossregion_v2"
OUTPUT = ROOT / "revision_v58_regional_consensus"
ANNOTATION = Path(os.environ["NMF_SOURCE_ROOT"]) / "tables" / "TableS3_program_annotation.tsv"

OUTPUT_TSV = OUTPUT / "data" / "TableS4_donor_aware_regional_programs.tsv"
OUTPUT_XLSX = OUTPUT / "data" / "TableS4_donor_aware_regional_programs.xlsx"

PROGRAMS = [f"P{i}" for i in range(1, 55)]
OVERALL_SOURCES = {
    "nature_lmm": ("overall_nature_lmm.tsv", "F_region", "p_region", "BH_54", "significant_BH_54"),
    "limma_dupcor": ("overall_limma_dupcor.tsv", "F_region", "p_region", "BH_54", "significant_BH_54"),
    "dream": ("overall_dream.tsv", "F", "P.Value", "BH_54", "significant_BH_54"),
}
SUBCLASS_SOURCES = {
    "nature_lmm": "subclass_nature_lmm.tsv",
    "limma_dupcor": "subclass_limma_dupcor.tsv",
    "dream": "subclass_dream.tsv",
}


def read_tsv(path: Path) -> pd.DataFrame:
    if not path.is_file():
        raise FileNotFoundError(path)
    return pd.read_csv(path, sep="\t", dtype=str, keep_default_na=False)


def as_bool(series: pd.Series, label: str) -> pd.Series:
    values = series.astype(str).str.strip().str.lower()
    allowed = {"true": True, "false": False, "1": True, "0": False}
    invalid = sorted(set(values) - set(allowed))
    if invalid:
        raise ValueError(f"{label} has invalid Boolean values: {invalid}")
    return values.map(allowed).astype(bool)


def check_program_set(df: pd.DataFrame, column: str, label: str) -> None:
    observed = list(df[column])
    if len(df) != 54 or set(observed) != set(PROGRAMS) or df[column].duplicated().any():
        raise ValueError(f"{label} must contain each of the 54 retained programs exactly once")


def program_order(series: pd.Series) -> pd.Series:
    return series.str.extract(r"(\d+)", expand=False).astype(int)


def add_check(rows: list[dict], check: str, observed: str, expected: str, passed: bool) -> None:
    rows.append(
        {
            "check": check,
            "observed": observed,
            "expected": expected,
            "status": "PASS" if passed else "FAIL",
        }
    )


def main() -> None:
    checks: list[dict] = []
    annotation = read_tsv(ANNOTATION)
    check_program_set(annotation, "new_P", "canonical TableS3_program_annotation")
    required_annotation = ["new_P", "functional_name", "dominant_class", "dominant_subclass"]
    missing_annotation = [column for column in required_annotation if column not in annotation.columns]
    if missing_annotation:
        raise ValueError("Missing canonical annotation columns: " + ", ".join(missing_annotation))
    annotation = annotation[required_annotation].rename(columns={"new_P": "program"})
    add_check(checks, "canonical_annotation_rows", str(len(annotation)), "54", len(annotation) == 54)
    add_check(checks, "canonical_annotation_unique_programs", str(annotation["program"].nunique()), "54", annotation["program"].nunique() == 54)

    table = annotation.copy()
    overall_flags: dict[str, pd.Series] = {}
    for method, (filename, f_column, p_column, bh_column, significant_column) in OVERALL_SOURCES.items():
        source = read_tsv(RESULTS / filename)
        check_program_set(source, "program", filename)
        required = ["program", f_column, p_column, bh_column, significant_column]
        missing = [column for column in required if column not in source.columns]
        if missing:
            raise ValueError(f"{filename} is missing columns: {', '.join(missing)}")
        source = source[required].copy()
        source[significant_column] = as_bool(source[significant_column], f"{filename}:{significant_column}")
        for column in [f_column, p_column, bh_column]:
            source[column] = pd.to_numeric(source[column], errors="raise")
        source = source.rename(
            columns={
                f_column: f"{method}_overall_F",
                p_column: f"{method}_overall_raw_P",
                bh_column: f"{method}_overall_BH_54",
                significant_column: f"{method}_overall_significant_BH_54",
            }
        )
        table = table.merge(source, on="program", how="left", validate="one_to_one")
        overall_flags[method] = table[f"{method}_overall_significant_BH_54"]
        add_check(checks, f"{method}_overall_rows", str(len(source)), "54", len(source) == 54)

    method_matrix = read_tsv(RESULTS / "program_method_matrix.tsv")
    check_program_set(method_matrix, "program", "program_method_matrix.tsv")
    consensus_source = read_tsv(RESULTS / "regional_consensus_21.tsv")
    if consensus_source["program"].duplicated().any() or len(consensus_source) != 21:
        raise ValueError("regional_consensus_21.tsv must contain 21 unique programs")
    if not set(consensus_source["program"]).issubset(set(PROGRAMS)):
        raise ValueError("regional_consensus_21.tsv includes a program outside P1-P54")

    matrix_methods = {
        "nature_lmm": "overall_nature_lmm",
        "limma_dupcor": "overall_limma_dupcor",
        "dream": "overall_dream",
    }
    source_intersection = pd.Series(True, index=table.index)
    for method, matrix_column in matrix_methods.items():
        if matrix_column not in method_matrix.columns:
            raise ValueError(f"program_method_matrix.tsv is missing {matrix_column}")
        matrix_values = pd.to_numeric(method_matrix.set_index("program").loc[table["program"], matrix_column], errors="raise").astype(int).to_numpy()
        source_values = table[f"{method}_overall_significant_BH_54"].astype(int).to_numpy()
        add_check(
            checks,
            f"{method}_overall_matches_method_matrix",
            str(int((matrix_values == source_values).sum())),
            "54 matching programs",
            bool((matrix_values == source_values).all()),
        )
        source_intersection &= table[f"{method}_overall_significant_BH_54"]

    table["three_model_consensus"] = table["program"].isin(set(consensus_source["program"])).map({True: "Yes", False: "No"})
    consensus_matches = set(table.loc[source_intersection, "program"]) == set(consensus_source["program"])
    add_check(checks, "three_model_intersection_matches_consensus_source", str(int(source_intersection.sum())), "21 exact programs", consensus_matches)
    add_check(checks, "three_model_consensus_programs", str(int((table["three_model_consensus"] == "Yes").sum())), "21", int((table["three_model_consensus"] == "Yes").sum()) == 21)

    for method, filename in SUBCLASS_SOURCES.items():
        source = read_tsv(RESULTS / filename)
        program_column = "program"
        significant_column = "significant_BH_global_22x54"
        if program_column not in source.columns or significant_column not in source.columns:
            raise ValueError(f"{filename} is missing {program_column} or {significant_column}")
        if len(source) != 54 * 22 or set(source[program_column]) != set(PROGRAMS):
            raise ValueError(f"{filename} must contain 54 programs across 22 subclasses")
        source[significant_column] = as_bool(source[significant_column], f"{filename}:{significant_column}")
        counts = source.groupby(program_column, sort=False)[significant_column].sum().reindex(PROGRAMS).astype(int)
        matrix_column = f"n_subclasses_{method}"
        if matrix_column not in method_matrix.columns:
            raise ValueError(f"program_method_matrix.tsv is missing {matrix_column}")
        matrix_counts = pd.to_numeric(method_matrix.set_index("program").loc[PROGRAMS, matrix_column], errors="raise").astype(int)
        add_check(checks, f"{method}_subclass_rows", str(len(source)), "1188", len(source) == 1188)
        add_check(
            checks,
            f"{method}_subclass_counts_match_method_matrix",
            str(int((counts.to_numpy() == matrix_counts.to_numpy()).sum())),
            "54 matching programs",
            bool((counts.to_numpy() == matrix_counts.to_numpy()).all()),
        )
        table[f"{method}_significant_subclass_count_global_BH_22x54"] = table["program"].map(counts)

    table = table.rename(
        columns={
            "functional_name": "program_functional_name",
            "dominant_class": "dominant_cell_class",
            "dominant_subclass": "dominant_cell_subclass",
        }
    )
    ordered_columns = [
        "program",
        "program_functional_name",
        "dominant_cell_class",
        "dominant_cell_subclass",
        "three_model_consensus",
        "nature_lmm_overall_F",
        "nature_lmm_overall_raw_P",
        "nature_lmm_overall_BH_54",
        "nature_lmm_overall_significant_BH_54",
        "limma_dupcor_overall_F",
        "limma_dupcor_overall_raw_P",
        "limma_dupcor_overall_BH_54",
        "limma_dupcor_overall_significant_BH_54",
        "dream_overall_F",
        "dream_overall_raw_P",
        "dream_overall_BH_54",
        "dream_overall_significant_BH_54",
        "nature_lmm_significant_subclass_count_global_BH_22x54",
        "limma_dupcor_significant_subclass_count_global_BH_22x54",
        "dream_significant_subclass_count_global_BH_22x54",
    ]
    table = table[ordered_columns].sort_values("program", key=program_order).reset_index(drop=True)
    add_check(checks, "tableS4_rows", str(len(table)), "54", len(table) == 54)
    add_check(checks, "tableS4_columns", str(len(table.columns)), str(len(ordered_columns)), len(table.columns) == len(ordered_columns))
    add_check(checks, "tableS4_unique_programs", str(table["program"].nunique()), "54", table["program"].nunique() == 54)
    add_check(checks, "tableS4_missing_values", str(int(table.isna().sum().sum())), "0", int(table.isna().sum().sum()) == 0)

    forbidden_tokens = ["old8", "eta", "pooled", "mouse", "internal", "note"]
    header_text = "|".join(table.columns).lower()
    present_forbidden = [token for token in forbidden_tokens if token in header_text]
    add_check(checks, "forbidden_legacy_or_internal_headers", ",".join(present_forbidden) or "none", "none", not present_forbidden)

    table.to_csv(OUTPUT_TSV, sep="\t", index=False, lineterminator="\n")
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="Table S4", index=False)
    workbook = load_workbook(OUTPUT_XLSX)
    worksheet = workbook["Table S4"]
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.zoomScale = 85
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for column_cells in worksheet.columns:
        letter = column_cells[0].column_letter
        longest = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        worksheet.column_dimensions[letter].width = min(max(longest + 2, 12), 40)
    for row in worksheet.iter_rows(min_row=2, min_col=6, max_col=17):
        for cell in row:
            cell.number_format = "0.000E+00"
    workbook.save(OUTPUT_XLSX)



if __name__ == "__main__":
    main()
