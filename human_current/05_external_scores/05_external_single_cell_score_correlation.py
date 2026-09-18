#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import gc
import io
import importlib.util
import json
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Iterable

import anndata as ad
import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from scipy import sparse

PROJECT = Path(os.environ["NMF_SOURCE_ROOT"]) / "inputs/cortex_nmf_program"
BENCH = PROJECT / "R2_Benchmark_Staging/01_human_single_cell_benchmark"
INPUT = BENCH / "input"
CORE = Path(__file__).resolve().parent
OUTPUT = Path((__import__("os").environ["NMF_WORK_ROOT"] + "/analysis/external_single_cell_program_scores"))
NEURON_OUTPUT = OUTPUT / "neuron_program_consistency_by_study.tsv"
CELL_NEAREST_OUTPUT_PREFIX = "cell_nearest_reference_"
CELL_NEAREST_BOX_PDF = OUTPUT / "cell_nearest_reference_cosine_boxplot.pdf"
CELL_NEAREST_BOX_PNG = OUTPUT / "cell_nearest_reference_cosine_boxplot.png"
CELL_NEAREST_QUERY_BATCH_SIZE = 4096
CELL_NEAREST_REFERENCE_BLOCK_SIZE = 262144
SCRIPT_OUTPUT = Path(__file__).resolve()

REFERENCE_PATH = PROJECT / "archived/results/cnmf_snrna_joint_full1M_v1/cnmf_work/snrna_joint_full1M_v1/snrna_joint_full1M_v1.starcat_spectra.k_60.dt_0_15.txt"
RETAIN_MAP_PATH = PROJECT / "archived/results/crossregion_v1/program_renumber_map.tsv"
DISCOVERY_PROFILE_PATH = PROJECT / "archived/results/crossregion_v1/cell_program_region_subclass.parquet"
DISCOVERY_OBS_PATH = PROJECT / "archived/inputs/snRNA_1M_obs.csv"
GENE_INFO_PATH = PROJECT / "archived/inputs/geneInfo_snRNA.csv"
SEAAD_META = PROJECT / "original_data/human/SEAAD_DLPFC_MTG"
SEAAD_ADAPTER_PATH = Path(__file__).with_name("00_scanvi_source_adapter.py")

STUDIES = [
    "Allen",
    "SEAAD",
    "Grubman_2019_GSE138852",
    "Lake_2018_GSE97930",
    "Schirmer_2019_MS_UCSC",
    "SingleSoma_AD_PFC_CELLxGENE",
    "Tran_2021_reward_cortex_CELLxGENE",
    "GSE144136",
    "GSE174367",
    "GSE291605",
]

BROAD_CONSISTENCY_STUDIES = [
    "Allen",
    "SEAAD",
    "Schirmer_2019_MS_UCSC",
    "SingleSoma_AD_PFC_CELLxGENE",
    "Tran_2021_reward_cortex_CELLxGENE",
    "GSE144136",
    "GSE174367",
    "GSE291605",
]
BROAD_CONSISTENCY_CLASSES = ["Exc", "Inh", "Ast", "Oligo", "OPC", "Micro", "Endo", "VLMC"]

ALLEN_H5AD = INPUT / "Allen/Human_Multiple_Cortical_Areas_SMARTseq/Human_Multiple_Cortical_Areas_SMART-seq__cellxgene_905763d1-e6fc-46d1-be5f-b1b232d14ce2.h5ad"
GRUBMAN_LIBS = [
    INPUT / "Grubman_2019_GSE138852/libraries/AD1_AD2.h5ad",
    INPUT / "Grubman_2019_GSE138852/libraries/AD3_AD4.h5ad",
    INPUT / "Grubman_2019_GSE138852/libraries/AD5_AD6.h5ad",
    INPUT / "Grubman_2019_GSE138852/libraries/Ct1_Ct2.h5ad",
    INPUT / "Grubman_2019_GSE138852/libraries/Ct3_Ct4.h5ad",
    INPUT / "Grubman_2019_GSE138852/libraries/Ct5_Ct6.h5ad",
]
LAKE_LIBS = [
    *[INPUT / f"Lake_2018_GSE97930/libraries/fcx{i}.h5ad" for i in range(1, 14)],
    *[INPUT / f"Lake_2018_GSE97930/libraries/occ{i}.h5ad" for i in range(1, 25)],
]
SCHIRMER_LIBS = [
    *[INPUT / f"Schirmer_2019_MS_UCSC/libraries/C{i}.h5ad" for i in range(1, 10)],
    *[INPUT / f"Schirmer_2019_MS_UCSC/libraries/MS{i}.h5ad" for i in range(1, 13)],
]
SINGLESOMA_H5AD = INPUT / "SingleSoma_AD_PFC_CELLxGENE/tangle_bearing_PFC_23197.h5ad"
TRAN_H5ADS = [
    ("DLPFC", INPUT / "Tran_2021_reward_cortex_CELLxGENE/DLPFC_11202.h5ad"),
    ("SACC", INPUT / "Tran_2021_reward_cortex_CELLxGENE/SACC_15669.h5ad"),
]
GENERIC_H5ADS = {
    "GSE144136": INPUT / "GSE144136/libraries/GSE144136_query.h5ad",
    "GSE174367": INPUT / "GSE174367/libraries/GSE174367_query.h5ad",
    "GSE291605": INPUT / "GSE291605/libraries/GSE291605_query.h5ad",
}
SEAAD_COUNTS = {
    ("normal", "dlpfc"): 99,
    ("normal", "mtg"): 105,
    ("dementia", "dlpfc"): 91,
    ("dementia", "mtg"): 101,
}

PROG60 = [str(i) for i in range(1, 61)]
GENE_INFO_CACHE: dict[str, str] | None = None


def clean(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().strip('"').strip("'")
    if text.lower() in {"nan", "none", "na", "<na>"}:
        return ""
    return text


def norm_gene(value: object) -> str:
    text = clean(value)
    if "|" in text:
        text = text.split("|")[-1]
    if text.startswith("ENSG") and "." in text:
        text = text.split(".", 1)[0]
    return text


def load_reference() -> tuple[np.ndarray, list[str], dict[int, int], list[str]]:
    spectra = pd.read_csv(REFERENCE_PATH, sep="\t", index_col=0)
    row_ids = [norm_gene(x) for x in spectra.index]
    row_number = {}
    for i, value in enumerate(row_ids):
        token = value.removeprefix("GEP").removeprefix("P")
        if token.isdigit():
            row_number[int(token)] = i
    if len(row_number) < 60:
        raise RuntimeError("K60 reference does not expose 60 numbered programs")
    ref_genes = [norm_gene(x) for x in spectra.columns]
    if len(ref_genes) != len(set(ref_genes)):
        raise RuntimeError("K60 reference gene axis is not unique")
    H = spectra.to_numpy(dtype=np.float64)
    return H, ref_genes, row_number, row_ids


def load_retained_map(row_number: dict[int, int]) -> tuple[list[int], list[str], pd.DataFrame]:
    table = pd.read_csv(RETAIN_MAP_PATH, sep="\t")
    required = {"old_P", "new_P", "status"}
    if not required.issubset(table.columns):
        raise RuntimeError("retained program map lacks old_P/new_P/status")
    keep = table[table["status"].astype(str).str.lower().eq("kept")].copy()
    keep["old_num"] = keep["old_P"].astype(str).str.replace("P", "", regex=False).astype(int)
    keep["new_num"] = keep["new_P"].astype(str).str.replace("P", "", regex=False).astype(int)
    keep = keep.sort_values("new_num", kind="stable")
    if keep["new_num"].tolist() != list(range(1, 55)):
        raise RuntimeError("retained program map is not the contiguous 54-program map")
    if any(x not in row_number for x in keep["old_num"].tolist()):
        raise RuntimeError("retained map references a missing K60 program")
    old_indices = [row_number[x] for x in keep["old_num"].tolist()]
    new_ids = [f"P{x}" for x in keep["new_num"].tolist()]
    return old_indices, new_ids, keep


def read_mapping_table(path: Path, header_prefix: str) -> pd.DataFrame:
    lines = path.read_text(encoding="utf-8").splitlines()
    while lines and not lines[0].startswith(header_prefix):
        lines.pop(0)
    if not lines:
        return pd.DataFrame()
    return pd.read_csv(io.StringIO("\n".join(lines)), sep="\t", dtype=str).fillna("")


def load_label_map(study: str, region: str | None = None) -> dict[str, str]:
    if study == "Allen":
        path = INPUT / "Allen/backup/label_mapping.tsv"
        table = read_mapping_table(path, "side\t")
        table = table[(table["side"] == "author") & table["common_label"].ne("")]
        table = table[table["mapping_status"] != "OUT_OF_SCOPE"]
        out: dict[str, str] = {}
        for row in table.itertuples(index=False):
            source = clean(row.source_label)
            target = clean(row.common_label)
            if not source or not target:
                continue
            if source in out and out[source] != target:
                out.pop(source, None)
            else:
                out[source] = target
        return out
    if study == "SEAAD":
        if region is None:
            raise ValueError("SEAAD label map requires DLPFC or MTG region")
        region_key = clean(region).upper()
        if region_key not in {"DLPFC", "MTG"}:
            raise ValueError(f"unsupported SEAAD region: {region}")
        paths = [INPUT / f"SEAAD/backup/SEAAD_{region_key}_Subclass_to_canonical22.tsv"]
        out = {}
        for path in paths:
            table = read_mapping_table(path, "source_label\t")
            for row in table.itertuples(index=False):
                source = clean(row.source_label)
                target = clean(row.canonical22)
                if not source or not target:
                    continue
                if source in out and out[source] != target:
                    out.pop(source, None)
                else:
                    out[source] = target
        return out
    path = INPUT / study / "backup/author_label_to_canonical22_crosswalk.tsv"
    table = read_mapping_table(path, "dataset\t")
    out = {}
    for row in table.itertuples(index=False):
        source = clean(getattr(row, "author_label", ""))
        target = clean(getattr(row, "canonical22", ""))
        status = clean(getattr(row, "mapping_status", ""))
        if not source or not target or status not in {"exact", "compatible-unique", "compatible-coarse"}:
            continue
        if source in out and out[source] != target:
            out.pop(source, None)
        else:
            out[source] = target
    return out


def map_label(raw: object, label_map: dict[str, str]) -> str:
    text = clean(raw)
    return label_map.get(text, "")


def common_label_for(study: str, author_label: object, region: object, label_maps) -> str:
    if study == "SEAAD":
        region_key = clean(region).upper()
        return map_label(author_label, label_maps.get(region_key, {}))
    return map_label(author_label, label_maps)


def sparse_counts(value: object) -> sparse.csr_matrix:
    if sparse.issparse(value):
        return value.tocsr().astype(np.float32, copy=False)
    return sparse.csr_matrix(np.asarray(value, dtype=np.float32))


def symbols_from_var(var: pd.DataFrame, names: Iterable[object]) -> list[str]:
    raw_names = [norm_gene(x) for x in names]
    if "feature_name" in var.columns:
        symbols = [clean(x) for x in var["feature_name"].tolist()]
        if len(symbols) == len(raw_names):
            return [norm_gene(x) for x in symbols]
    if any(x.startswith("ENSG") for x in raw_names):
        gene_info = load_gene_info()
        return [gene_info.get(x, x) for x in raw_names]
    return raw_names


def h5ad_unit(
    path: Path,
    unit_name: str,
    *,
    use_raw: bool,
    label_field: str,
    donor_field: str,
    region_field: str,
    condition_field: str,
    label_override: pd.Series | None = None,
    condition_override: pd.Series | None = None,
) -> dict[str, object]:
    source = ad.read_h5ad(path)
    obs = source.obs.copy()
    cells = np.asarray([clean(x) for x in source.obs_names], dtype=str)
    if use_raw and source.raw is not None:
        matrix = sparse_counts(source.raw.X)
        genes = symbols_from_var(source.raw.var, source.raw.var_names)
    else:
        matrix = sparse_counts(source.X)
        genes = symbols_from_var(source.var, source.var_names)
    if matrix.shape[0] != len(obs):
        raise RuntimeError(f"{path}: matrix/obs row mismatch")
    def field(name: str, fallback: str = "") -> np.ndarray:
        if name and name in obs.columns:
            return np.asarray([clean(x) for x in obs[name].tolist()], dtype=str)
        return np.repeat(fallback, len(obs)).astype(str)
    labels = field(label_field)
    if label_override is not None:
        labels = np.asarray([clean(label_override.get(x, "")) for x in cells], dtype=str)
    conditions = field(condition_field)
    if condition_override is not None:
        conditions = np.asarray([clean(condition_override.get(x, "")) for x in cells], dtype=str)
    donors = field(donor_field, unit_name)
    regions = field(region_field, unit_name)
    source_rows = (
        pd.to_numeric(obs["source_row"], errors="coerce").fillna(-1).astype(np.int64).to_numpy()
        if "source_row" in obs.columns else np.arange(len(obs), dtype=np.int64)
    )
    meta = pd.DataFrame({
        "cell_id": cells,
        "source_row": source_rows,
        "source_unit": np.repeat(unit_name, len(obs)),
        "donor": donors,
        "region": regions,
        "condition": conditions,
        "author_label": labels,
    })
    del source, obs
    gc.collect()
    return {"name": unit_name, "X": matrix, "genes": genes, "meta": meta}


def load_grubman_annotations() -> tuple[pd.DataFrame, pd.Series, pd.Series]:
    root = INPUT / "Grubman_2019_GSE138852"
    cov = pd.read_csv(root / "GSE138852_covariates.csv.gz", compression="gzip", index_col=0, dtype=str).fillna("")
    return cov, cov["oupSample.cellType"], cov["oupSample.batchCond"]


def load_lake_annotations() -> dict[str, dict[str, str]]:
    from openpyxl import load_workbook
    path = INPUT / "Lake_2018_GSE97930/Lake_2018_Supplementary_Tables_1_13.xlsx"
    book = load_workbook(path, read_only=True, data_only=True)
    out: dict[str, dict[str, str]] = {}
    for row in book["Table S2"].iter_rows(min_row=7, values_only=True):
        sample = row[1]
        if sample in (None, ""):
            continue
        identity = "" if row[6] is None else str(row[6])
        key = clean(f"{identity}_{sample}")
        out[key] = {
            "library": "" if row[2] is None else clean(row[2]),
            "identity": identity,
            "patient": "" if row[4] is None else clean(row[4]),
            "region": "" if row[5] is None else clean(row[5]),
        }
    return out


def load_schirmer_meta() -> pd.DataFrame:
    path = INPUT / "Schirmer_2019_MS_UCSC/meta.tsv"
    return pd.read_csv(path, sep="\t", dtype=str).fillna("").set_index("cell", drop=False)


def external_units(study: str, maps: dict[str, dict[str, str]]) -> Iterable[dict[str, object]]:
    if study == "Allen":
        yield h5ad_unit(
            ALLEN_H5AD, "whole_query", use_raw=False,
            label_field="subclass", donor_field="donor_id", region_field="region", condition_field="disease",
        )
        return
    if study == "Grubman_2019_GSE138852":
        cov, labels, conditions = load_grubman_annotations()
        for path in GRUBMAN_LIBS:
            yield h5ad_unit(
                path, path.stem, use_raw=False,
                label_field="", donor_field="donor", region_field="region", condition_field="",
                label_override=labels, condition_override=conditions,
            )
        return
    if study == "Lake_2018_GSE97930":
        annot = load_lake_annotations()
        for path in LAKE_LIBS:
            unit = h5ad_unit(
                path, path.stem, use_raw=False,
                label_field="", donor_field="donor", region_field="region", condition_field="",
            )
            cells = unit["meta"]["cell_id"].astype(str)
            labels = np.asarray([annot.get(clean(x), {}).get("identity", "") for x in cells], dtype=str)
            unit["meta"]["author_label"] = labels
            yield unit
        return
    if study == "Schirmer_2019_MS_UCSC":
        meta = load_schirmer_meta()
        for path in SCHIRMER_LIBS:
            unit = h5ad_unit(
                path, path.stem, use_raw=False,
                label_field="", donor_field="donor", region_field="region", condition_field="",
            )
            cells = unit["meta"]["cell_id"].astype(str)
            unit["meta"]["author_label"] = np.asarray([clean(meta.loc[x, "cell_type"]) if x in meta.index else "" for x in cells], dtype=str)
            unit["meta"]["condition"] = np.asarray([clean(meta.loc[x, "diagnosis"]) if x in meta.index else "" for x in cells], dtype=str)
            unit["meta"]["donor"] = np.asarray([clean(meta.loc[x, "sample"]) if x in meta.index else unit["name"] for x in cells], dtype=str)
            unit["meta"]["region"] = np.asarray([clean(meta.loc[x, "region"]) if x in meta.index else "" for x in cells], dtype=str)
            yield unit
        return
    if study == "SingleSoma_AD_PFC_CELLxGENE":
        yield h5ad_unit(
            SINGLESOMA_H5AD, "whole_query", use_raw=True,
            label_field="Cell.Types", donor_field="donor_id", region_field="tissue", condition_field="disease",
        )
        return
    if study == "Tran_2021_reward_cortex_CELLxGENE":
        for name, path in TRAN_H5ADS:
            yield h5ad_unit(
                path, name, use_raw=False,
                label_field="author_cell_type", donor_field="donor_id", region_field="tissue", condition_field="disease",
            )
        return
    if study in GENERIC_H5ADS:
        yield h5ad_unit(
            GENERIC_H5ADS[study], "whole_query", use_raw=False,
            label_field="author_label", donor_field="sample_id", region_field="region", condition_field="condition",
        )
        return
    raise RuntimeError(f"unsupported external study: {study}")


def load_seaad_adapter(cohort: str):
    cohort = clean(cohort).lower()
    if cohort not in {"normal", "dementia"}:
        raise ValueError(f"unsupported SEAAD cohort: {cohort}")
    previous = os.environ.get("SEAAD_COHORT")
    os.environ["SEAAD_COHORT"] = cohort
    try:
        spec = importlib.util.spec_from_file_location(
            f"seaad_source_adapter_{cohort}", SEAAD_ADAPTER_PATH
        )
        module = importlib.util.module_from_spec(spec)
        if spec.loader is None:
            raise RuntimeError("SEAAD adapter loader is unavailable")
        spec.loader.exec_module(module)
        return module
    finally:
        if previous is None:
            os.environ.pop("SEAAD_COHORT", None)
        else:
            os.environ["SEAAD_COHORT"] = previous


def rds_gene_axis(adapter, cohort: str, region: str) -> np.ndarray:
    from rpy2 import robjects as ro
    path = adapter.library_path(region, 1)
    obj = ro.r["readRDS"](str(path))
    counts = obj.rx2("counts")
    genes = np.asarray(list(counts.slots["Dimnames"])[0], dtype=str)
    del obj, counts
    gc.collect()
    return genes


def load_gene_info() -> dict[str, str]:
    global GENE_INFO_CACHE
    if GENE_INFO_CACHE is not None:
        return GENE_INFO_CACHE
    table = pd.read_csv(GENE_INFO_PATH, usecols=["gene_id", "gene_name"], dtype=str).fillna("")
    out = {}
    for row in table.itertuples(index=False):
        gid = norm_gene(row.gene_id)
        name = norm_gene(row.gene_name)
        if gid and name:
            out[gid] = name
    GENE_INFO_CACHE = out
    return GENE_INFO_CACHE


def seaad_obs(region: str) -> pd.DataFrame:
    path = SEAAD_META / f"SEAAD_{region.upper()}_reconstructed_raw_counts.h5ad"
    source = ad.read_h5ad(path, backed="r")
    obs = source.obs[["cell_id", "source_row", "Subclass", "donor_id", "region", "disease", "Specimen ID"]].copy()
    source.file.close()
    for col in obs.columns:
        obs[col] = obs[col].astype(str)
    obs["key"] = obs["cell_id"].map(clean) + "\x1f" + pd.to_numeric(obs["source_row"], errors="coerce").fillna(-1).astype(np.int64).astype(str)
    return obs.set_index("key", drop=False)


def seaad_units(adapter, cohort: str, region: str, gene_info: dict[str, str], label_map: dict[str, str]) -> Iterable[dict[str, object]]:
    genes_raw = rds_gene_axis(adapter, cohort, region)
    genes = [gene_info.get(norm_gene(x), norm_gene(x)) for x in genes_raw]
    obs_lookup = seaad_obs(region)
    count = SEAAD_COUNTS[(cohort, region)]
    for ordinal in range(1, count + 1):
        obj = adapter.load_library(region, ordinal, genes_raw)
        matrix = sparse_counts(obj.X)
        base = obj.obs.copy()
        cells = np.asarray([clean(x) for x in base["cell_id"].tolist()], dtype=str)
        source_rows = pd.to_numeric(base["source_row"], errors="coerce").fillna(-1).astype(np.int64).to_numpy()
        keys = [f"{c}\x1f{r}" for c, r in zip(cells, source_rows)]
        matched = obs_lookup.reindex(keys)
        labels = np.asarray([clean(x) for x in matched["Subclass"].tolist()], dtype=str)
        conditions = np.asarray([clean(x) for x in matched["disease"].tolist()], dtype=str)
        donors = np.asarray([clean(x) for x in base["donor"].tolist()], dtype=str)
        regions = np.asarray([clean(x) for x in base["region"].tolist()], dtype=str)
        units_meta = pd.DataFrame({
            "cell_id": cells,
            "source_row": source_rows,
            "source_unit": np.repeat(f"{cohort}_{region}_library_{ordinal:05d}", len(cells)),
            "donor": donors,
            "region": regions,
            "condition": conditions,
            "author_label": labels,
        })
        del obj, base, matched
        gc.collect()
        yield {"name": f"{cohort}_{region}_library_{ordinal:05d}", "X": matrix, "genes": genes, "meta": units_meta}


def all_seaad_units(adapters, gene_info: dict[str, str], label_maps) -> Iterable[dict[str, object]]:
    for cohort, region in (("normal", "dlpfc"), ("normal", "mtg"), ("dementia", "dlpfc"), ("dementia", "mtg")):
        yield from seaad_units(adapters[cohort], cohort, region, gene_info, label_maps[region.upper()])


def align_to_reference(X: sparse.csr_matrix, genes: list[str], ref_genes: list[str]) -> tuple[sparse.csr_matrix, list[str], np.ndarray]:
    ref_pos = {gene: i for i, gene in enumerate(ref_genes)}
    grouped: dict[int, list[int]] = {}
    for j, gene in enumerate(genes):
        gene = norm_gene(gene)
        if gene in ref_pos:
            grouped.setdefault(ref_pos[gene], []).append(j)
    if not grouped:
        raise RuntimeError("query has no overlap with the K60 reference")
    ref_idx = np.asarray(sorted(grouped), dtype=np.int64)
    query_positions: list[int] = []
    output_positions: list[int] = []
    for output_index, index in enumerate(ref_idx.tolist()):
        positions = grouped[index]
        query_positions.extend(positions)
        output_positions.extend([output_index] * len(positions))
    query_idx = np.asarray(query_positions, dtype=np.int64)
    subset = X[:, query_idx]
    if all(len(grouped[index]) == 1 for index in ref_idx.tolist()):
        aligned = subset.astype(np.float32, copy=False).tocsr()
    else:
        mapping = sparse.csr_matrix(
            (
                np.ones(len(query_idx), dtype=np.float32),
                (np.arange(len(query_idx), dtype=np.int64), np.asarray(output_positions, dtype=np.int64)),
            ),
            shape=(len(query_idx), len(ref_idx)),
        )
        aligned = (subset @ mapping).astype(np.float32, copy=False).tocsr()
        aligned.sum_duplicates()
    q_genes = [ref_genes[i] for i in ref_idx.tolist()]
    return aligned, q_genes, ref_idx


def study_units(study: str, label_map: dict[str, str]) -> Iterable[dict[str, object]]:
    if study == "SEAAD":
        adapters = {
            cohort: load_seaad_adapter(cohort) for cohort in ("normal", "dementia")
        }
        gene_info = load_gene_info()
        label_maps = {region: load_label_map("SEAAD", region) for region in ("DLPFC", "MTG")}
        yield from all_seaad_units(adapters, gene_info, label_maps)
    else:
        yield from external_units(study, {study: label_map})


def compute_study_sd(study: str, ref_genes: list[str], label_map: dict[str, str]) -> tuple[np.ndarray, int]:
    sums = np.zeros(len(ref_genes), dtype=np.float64)
    squares = np.zeros(len(ref_genes), dtype=np.float64)
    n_rows = 0
    for unit in study_units(study, label_map):
        X = sparse_counts(unit["X"])
        aligned, _, ref_idx = align_to_reference(X, [norm_gene(x) for x in unit["genes"]], ref_genes)
        sums[ref_idx] += np.asarray(aligned.sum(axis=0)).ravel()
        squares[ref_idx] += np.asarray(aligned.multiply(aligned).sum(axis=0)).ravel()
        n_rows += aligned.shape[0]
        del unit, X, aligned
        gc.collect()
    if n_rows <= 0:
        raise RuntimeError(f"{study}: no query rows")
    mean = sums / float(n_rows)
    variance = np.maximum(squares / float(n_rows) - mean * mean, 0.0)
    return np.sqrt(variance), n_rows


def normalized_selected(result: object, old_indices: list[int]) -> tuple[np.ndarray, np.ndarray]:
    full = np.asarray(result.normalized_usage, dtype=np.float32)
    selected = full[:, old_indices]
    mass = selected.sum(axis=1).astype(np.float32)
    return selected, mass


def write_parquet_stream(path: Path, frames: Iterable[pd.DataFrame]) -> int:
    writer = None
    rows = 0
    try:
        for frame in frames:
            table = pa.Table.from_pandas(frame, preserve_index=False)
            if writer is None:
                writer = pq.ParquetWriter(path, table.schema, compression="zstd")
            writer.write_table(table)
            rows += len(frame)
    finally:
        if writer is not None:
            writer.close()
    return rows


def aggregate_profile(
    accumulator: dict[tuple[str, str, str, str, str], tuple[np.ndarray, int]],
    meta: pd.DataFrame,
    scores: np.ndarray,
) -> None:
    keys = list(zip(
        meta["donor"].astype(str),
        meta["author_label"].astype(str),
        meta["common_label"].astype(str),
        meta["region"].astype(str),
        meta["condition"].astype(str),
    ))
    local: dict[tuple[str, str, str, str, str], tuple[np.ndarray, int]] = {}
    for i, key in enumerate(keys):
        if key in local:
            local[key] = (local[key][0] + scores[i], local[key][1] + 1)
        else:
            local[key] = (scores[i].astype(np.float64), 1)
    for key, (total, count) in local.items():
        if key in accumulator:
            accumulator[key] = (accumulator[key][0] + total, accumulator[key][1] + count)
        else:
            accumulator[key] = (total, count)


def accumulator_frame(
    study: str,
    accumulator: dict[tuple[str, str, str, str, str], tuple[np.ndarray, int]],
    program_names: list[str],
) -> pd.DataFrame:
    rows = []
    for (donor, label, common, region, condition), (total, count) in accumulator.items():
        row = {
            "study": study,
            "donor": donor,
            "author_label": label,
            "common_label": common,
            "region": region,
            "condition": condition,
            "n_cells": int(count),
        }
        row.update({name: float(value) for name, value in zip(program_names, total / float(count))})
        rows.append(row)
    columns = ["study", "donor", "author_label", "common_label", "region", "condition", "n_cells", *program_names]
    return pd.DataFrame(rows, columns=columns)


def refresh_existing_study(
    study: str,
    score_path: Path,
    profile_path: Path,
    label_map: dict[str, str],
    program_names: list[str],
) -> pd.DataFrame:
    frame = pd.read_parquet(score_path)
    frame["common_label"] = [
        common_label_for(study, label, region, label_map)
        for label, region in zip(frame["author_label"], frame["region"])
    ]
    frame.to_parquet(score_path, index=False)
    accumulator: dict[tuple[str, str, str, str, str], tuple[np.ndarray, int]] = {}
    aggregate_profile(accumulator, frame, frame[program_names].to_numpy(dtype=np.float32))
    profile = accumulator_frame(study, accumulator, program_names)
    profile.to_parquet(profile_path, index=False)
    return profile


def score_study(
    study: str,
    H: np.ndarray,
    ref_genes: list[str],
    old_indices: list[int],
    program_names: list[str],
    label_map: dict[str, str],
) -> tuple[Path, Path, pd.DataFrame]:
    import gep_score as gep_module

    sd_ref, n_rows = compute_study_sd(study, ref_genes, label_map)
    score_path = OUTPUT / f"{study}.cell_scores.parquet"
    profile_path = OUTPUT / f"{study}.group_profiles.parquet"
    accumulator: dict[tuple[str, str, str, str, str], tuple[np.ndarray, int]] = {}
    original_gene_sd = gep_module._gene_sd
    functional_indices = np.arange(H.shape[0], dtype=np.int64)

    def frames() -> Iterable[pd.DataFrame]:
        nonlocal accumulator
        for unit in study_units(study, label_map):
            X = sparse_counts(unit["X"])
            genes = [norm_gene(x) for x in unit["genes"]]
            aligned, query_genes, ref_idx = align_to_reference(X, genes, ref_genes)
            sd_sub = sd_ref[ref_idx].copy()
            gep_module._gene_sd = lambda _counts, values=sd_sub: values
            try:
                result = gep_module.score_counts_fixed_h(
                    aligned,
                    query_genes,
                    H,
                    ref_genes,
                    identity_indices=np.asarray([], dtype=np.int64),
                    functional_indices=functional_indices,
                    device="cuda",
                    dtype="float32",
                )
            finally:
                gep_module._gene_sd = original_gene_sd
            selected, mass = normalized_selected(result, old_indices)
            meta = unit["meta"].copy()
            meta["common_label"] = [
                common_label_for(study, label, region, label_map)
                for label, region in zip(meta["author_label"], meta["region"])
            ]
            aggregate_profile(accumulator, meta, selected)
            out = meta[["cell_id", "source_row", "source_unit", "donor", "region", "condition", "author_label", "common_label"]].copy()
            out["selected_usage_mass"] = mass
            for j, name in enumerate(program_names):
                out[name] = selected[:, j]
            yield out
            del unit, X, aligned, result, selected, mass, meta, out
            gc.collect()

    rows = write_parquet_stream(score_path, frames())
    profiles = accumulator_frame(study, accumulator, program_names)
    profiles.to_parquet(profile_path, index=False)
    status = pd.DataFrame([{
        "study": study,
        "status": "COMPLETE",
        "query_rows": int(rows),
        "study_rows_used_for_sd": int(n_rows),
        "score_path": str(score_path),
        "profile_path": str(profile_path),
    }])
    return score_path, profile_path, status


def discovery_profiles(
    H: np.ndarray,
    old_indices: list[int],
    program_names: list[str],
    label_map: dict[str, str],
) -> pd.DataFrame:
    profile = pd.read_parquet(DISCOVERY_PROFILE_PATH)
    usage = profile[PROG60].to_numpy(dtype=np.float64)
    denom = usage.sum(axis=1)
    normalized = np.divide(usage, denom[:, None], out=np.zeros_like(usage), where=denom[:, None] > 0)
    selected = normalized[:, old_indices]
    if "__index_level_0__" in profile.columns:
        cell_ids = np.asarray([clean(x) for x in profile["__index_level_0__"].tolist()], dtype=str)
    else:
        cell_ids = np.asarray([clean(x) for x in profile.index.tolist()], dtype=str)
    obs = pd.read_csv(DISCOVERY_OBS_PATH, index_col=0, dtype=str).fillna("")
    if len(obs) != len(profile):
        raise RuntimeError("discovery metadata/profile row counts differ")
    if not np.array_equal(obs.index.astype(str).to_numpy(), cell_ids):
        obs = obs.reindex(cell_ids)
    meta = pd.DataFrame({
        "donor": obs["donor"].astype(str).map(clean).to_numpy(),
        "author_label": obs["subclass"].astype(str).map(clean).to_numpy(),
        "region": obs["region"].astype(str).map(clean).to_numpy(),
        "condition": np.repeat("", len(obs)),
    })
    meta["common_label"] = meta["author_label"].astype(str).map(clean).to_numpy()
    accumulator: dict[tuple[str, str, str, str, str], tuple[np.ndarray, int]] = {}
    aggregate_profile(accumulator, meta, selected)
    result = accumulator_frame("discovery", accumulator, program_names)
    result.to_parquet(OUTPUT / "discovery.group_profiles.parquet", index=False)
    return result


def spearman_only(a: np.ndarray, b: np.ndarray) -> float:
    x = np.asarray(a, dtype=np.float64)
    y = np.asarray(b, dtype=np.float64)
    if x.size != y.size or x.size < 2 or not np.isfinite(x).all() or not np.isfinite(y).all():
        return float("nan")
    rx = pd.Series(x).rank(method="average").to_numpy(dtype=np.float64)
    ry = pd.Series(y).rank(method="average").to_numpy(dtype=np.float64)
    sx = float(rx.std())
    sy = float(ry.std())
    if sx == 0.0 or sy == 0.0:
        return float("nan")
    return float(np.mean((rx - rx.mean()) * (ry - ry.mean())) / (sx * sy))


def donor_group_profiles(
    group_profiles: dict[str, pd.DataFrame], program_names: list[str]
) -> dict[str, pd.DataFrame]:
    columns = [
        "study",
        "donor",
        "author_label",
        "common_label",
        "region",
        "condition",
        "n_cells",
        *program_names,
    ]
    out: dict[str, pd.DataFrame] = {}
    for study, frame in group_profiles.items():
        result = frame.copy()
        for key in ("donor", "author_label", "common_label", "region", "condition"):
            result[key] = result[key].map(clean)
        result["n_cells"] = pd.to_numeric(result["n_cells"], errors="coerce").fillna(0).astype(np.int64)
        out[study] = result[columns].reset_index(drop=True)
    return out


def rank_standardized(values: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    array = np.asarray(values, dtype=np.float64)
    finite = np.isfinite(array).all(axis=1)
    ranks = pd.DataFrame(array).rank(axis=1, method="average").to_numpy(dtype=np.float64)
    centered = ranks - ranks.mean(axis=1, keepdims=True)
    norm = np.sqrt(np.square(centered).sum(axis=1))
    valid = finite & np.isfinite(norm) & (norm > 0.0)
    standardized = np.zeros_like(centered)
    np.divide(centered, norm[:, None], out=standardized, where=norm[:, None] > 0.0)
    return standardized, valid


CORRELATION_COLUMNS = [
    "study",
    "discovery_donor",
    "discovery_author_label",
    "discovery_common_label",
    "discovery_region",
    "discovery_condition",
    "discovery_n_cells",
    "external_donor",
    "external_author_label",
    "external_common_label",
    "external_region",
    "external_condition",
    "external_n_cells",
    "same_author_celltype",
    "status",
    "spearman_rho",
]


POSTPROCESS_STUDIES = [
    "Allen",
    "SEAAD",
    "Grubman_2019_GSE138852",
    "Lake_2018_GSE97930",
    "Schirmer_2019_MS_UCSC",
    "SingleSoma_AD_PFC_CELLxGENE",
    "Tran_2021_reward_cortex_CELLxGENE",
    "GSE144136",
    "GSE174367",
    "GSE291605",
]


MERGE_STUDIES = [
    "Allen",
    "SEAAD",
    "Grubman_2019_GSE138852",
    "Lake_2018_GSE97930",
    "Schirmer_2019_MS_UCSC",
    "SingleSoma_AD_PFC_CELLxGENE",
    "Tran_2021_reward_cortex_CELLxGENE",
    "GSE144136",
    "GSE174367",
    "GSE291605",
]


def paired_correlation_frame(
    study: str,
    discovery: pd.DataFrame,
    external: pd.DataFrame,
    discovery_z: np.ndarray,
    discovery_valid: np.ndarray,
    external_z: np.ndarray,
    external_valid: np.ndarray,
) -> pd.DataFrame:
    n_discovery = len(discovery)
    n_external = len(external)
    discovery_index = np.tile(np.arange(n_discovery, dtype=np.int64), n_external)
    external_index = np.repeat(np.arange(n_external, dtype=np.int64), n_discovery)
    rho = (discovery_z @ external_z.T).T.reshape(-1)
    valid = discovery_valid[discovery_index] & external_valid[external_index]
    rho = rho.astype(np.float64, copy=False)
    rho[~valid] = np.nan
    discovery_common = discovery["common_label"].to_numpy(dtype=str)
    external_common = external["common_label"].to_numpy(dtype=str)
    same = (
        (discovery_common[discovery_index] == external_common[external_index])
        & (discovery_common[discovery_index] != "")
        & (external_common[external_index] != "")
    )
    status = np.where(same, "MATCHED_AUTHOR_CELLTYPE", "NONCORRESPONDING_TYPE_REFERENCE").astype(object)
    status[~valid] = np.asarray(
        [f"{value}:CONSTANT_OR_MISSING_PROFILE" for value in status[~valid]], dtype=object
    )
    d = discovery.iloc[discovery_index].reset_index(drop=True)
    e = external.iloc[external_index].reset_index(drop=True)
    return pd.DataFrame(
        {
            "study": np.repeat(study, len(discovery_index)),
            "discovery_donor": d["donor"].to_numpy(dtype=str),
            "discovery_author_label": d["author_label"].to_numpy(dtype=str),
            "discovery_common_label": d["common_label"].to_numpy(dtype=str),
            "discovery_region": d["region"].to_numpy(dtype=str),
            "discovery_condition": d["condition"].to_numpy(dtype=str),
            "discovery_n_cells": d["n_cells"].to_numpy(dtype=np.int64),
            "external_donor": e["donor"].to_numpy(dtype=str),
            "external_author_label": e["author_label"].to_numpy(dtype=str),
            "external_common_label": e["common_label"].to_numpy(dtype=str),
            "external_region": e["region"].to_numpy(dtype=str),
            "external_condition": e["condition"].to_numpy(dtype=str),
            "external_n_cells": e["n_cells"].to_numpy(dtype=np.int64),
            "same_author_celltype": same,
            "status": status,
            "spearman_rho": rho,
        },
        columns=CORRELATION_COLUMNS,
    )


def unmapped_correlation_frame(study: str, external: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for row in external.itertuples(index=False):
        rows.append(
            {
                "study": study,
                "discovery_donor": "",
                "discovery_author_label": "",
                "discovery_common_label": "",
                "discovery_region": "",
                "discovery_condition": "",
                "discovery_n_cells": 0,
                "external_donor": clean(row.donor),
                "external_author_label": clean(row.author_label),
                "external_common_label": "",
                "external_region": clean(row.region),
                "external_condition": clean(row.condition),
                "external_n_cells": int(row.n_cells),
                "same_author_celltype": False,
                "status": "UNMAPPED_AUTHOR_LABEL",
                "spearman_rho": np.nan,
            }
        )
    return pd.DataFrame(rows, columns=CORRELATION_COLUMNS)


def write_ranked_correlations(full_path: Path, ranked_path: Path) -> None:
    fd_key, key_name = tempfile.mkstemp(prefix="external_corr_sort_", dir="/tmp")
    os.close(fd_key)
    fd_sorted, sorted_name = tempfile.mkstemp(prefix="external_corr_sorted_", dir="/tmp")
    os.close(fd_sorted)
    key_path = Path(key_name)
    sorted_path = Path(sorted_name)
    try:
        with key_path.open("w", encoding="utf-8") as keyed:
            subprocess.run(
                [
                    "awk",
                    "-F",
                    "\t",
                    'NR==1{next} {k=$16; if(k=="" || k=="nan" || k=="NaN") k="-inf"; printf "%s\\t%s\\n", k, $0}',
                    str(full_path),
                ],
                check=True,
                stdout=keyed,
            )
        with sorted_path.open("w", encoding="utf-8") as sorted_out:
            subprocess.run(
                ["sort", "-t", "\t", "-k1,1gr", str(key_path)],
                check=True,
                stdout=sorted_out,
            )
        with full_path.open("r", encoding="utf-8") as source:
            header = source.readline()
        with ranked_path.open("w", encoding="utf-8") as ranked:
            ranked.write(header)
            ranked.flush()
            subprocess.run(
                ["cut", "-f", "2-", str(sorted_path)],
                check=True,
                stdout=ranked,
            )
    finally:
        key_path.unlink(missing_ok=True)
        sorted_path.unlink(missing_ok=True)


def _label_maps_for_study(study: str):
    if study == "SEAAD":
        return {
            region: load_label_map("SEAAD", region)
            for region in ("DLPFC", "MTG")
        }
    return load_label_map(study)


def _apply_common_labels(study: str, frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    label_maps = _label_maps_for_study(study)
    result["common_label"] = [
        common_label_for(study, label, region, label_maps)
        for label, region in zip(result["author_label"], result["region"])
    ]
    return result


def postprocess_one_study(study: str) -> None:
    """Write only the three stratified correlation tables for one study."""
    if study not in POSTPROCESS_STUDIES:
        raise ValueError(f"unsupported post-processing study: {study}")

    program_names = [f"P{i}" for i in range(1, 55)]
    discovery = pd.read_parquet(OUTPUT / "discovery.group_profiles.parquet")
    external_path = OUTPUT / f"{study}.group_profiles.parquet"
    external = pd.read_parquet(external_path)

    external = _apply_common_labels(study, external)

    profile_units = donor_group_profiles(
        {"discovery": discovery, study: external}, program_names
    )
    discovery_units = profile_units["discovery"]
    external_units = profile_units[study]
    discovery_z, discovery_valid = rank_standardized(
        discovery_units[program_names].to_numpy(dtype=np.float64)
    )
    mapped = external_units[external_units["common_label"].ne("")].reset_index(drop=True)
    unmapped = external_units[external_units["common_label"].eq("")].reset_index(drop=True)

    full_frames = []
    for start in range(0, len(mapped), 256):
        chunk = mapped.iloc[start : start + 256].reset_index(drop=True)
        external_z, external_valid = rank_standardized(
            chunk[program_names].to_numpy(dtype=np.float64)
        )
        full_frames.append(
            paired_correlation_frame(
                study,
                discovery_units,
                chunk,
                discovery_z,
                discovery_valid,
                external_z,
                external_valid,
            )
        )
    if len(unmapped):
        full_frames.append(unmapped_correlation_frame(study, unmapped))

    full_path = OUTPUT / f"{study}.profile_correlations.tsv"
    same_path = OUTPUT / f"{study}.profile_correlations_same_author_celltype.tsv"
    ranked_path = OUTPUT / f"{study}.profile_correlations_ranked.tsv"
    if full_frames:
        full = pd.concat(full_frames, ignore_index=True)
    else:
        full = pd.DataFrame(columns=CORRELATION_COLUMNS)
    full.to_csv(full_path, sep="\t", index=False)
    full.loc[full["same_author_celltype"].eq(True)].to_csv(
        same_path, sep="\t", index=False
    )
    write_ranked_correlations(full_path, ranked_path)


def merge_all_outputs() -> None:
    full_paths = [
        OUTPUT / f"{study}.profile_correlations.tsv"
        for study in MERGE_STUDIES
    ]
    same_paths = [
        OUTPUT / f"{study}.profile_correlations_same_author_celltype.tsv"
        for study in MERGE_STUDIES
    ]
    public_full = OUTPUT / "external_profile_correlations.tsv"
    public_same = OUTPUT / "external_profile_correlations_same_author_celltype.tsv"
    public_ranked = OUTPUT / "external_profile_correlations_ranked.tsv"

    for destination, sources in ((public_full, full_paths), (public_same, same_paths)):
        with destination.open("w", encoding="utf-8") as target:
            for index, source in enumerate(sources):
                with source.open("r", encoding="utf-8") as handle:
                    header = handle.readline()
                    if index == 0:
                        target.write(header)
                    for chunk in iter(lambda: handle.read(1024 * 1024), ""):
                        target.write(chunk)
    write_ranked_correlations(public_full, public_ranked)

    for study in ["discovery", *MERGE_STUDIES]:
        profile_path = OUTPUT / f"{study}.group_profiles.parquet"
        profile = pd.read_parquet(profile_path)
        if study != "discovery":
            profile = _apply_common_labels(study, profile)
        profile.to_csv(OUTPUT / f"{study}.label_profiles.tsv", sep="\t", index=False)


UMAP_PROGRAMS = [f"P{i}" for i in range(1, 55)]
UMAP_PROFILE_COLUMNS = [
    "point_id",
    "point_kind",
    "study",
    "library",
    "source_unit",
    "donor",
    "author_label",
    "common_label",
    "mapping_status",
    "mapping_source",
    "region",
    "condition",
    "n_cells",
    "unit_kind",
    "library_source",
    "library_is_proxy",
    "library_source_status",
    "profile_status",
    *UMAP_PROGRAMS,
]

UMAP_PLOT_CLASS_ORDER = [
    "EXC",
    "INH",
    "OLIGO",
    "OPC",
    "MICRO",
    "AST",
    "ENDO",
    "VLMC",
    "Unmapped",
]
UMAP_PLOT_CLASS_COLORS = {
    "EXC": "#1f77b4",
    "INH": "#d62728",
    "OLIGO": "#2ca02c",
    "OPC": "#9467bd",
    "MICRO": "#8c564b",
    "AST": "#ff7f0e",
    "ENDO": "#17becf",
    "VLMC": "#e377c2",
    "Unmapped": "#7f7f7f",
}
UMAP_ALLOWED_PLOT_CLASSES = set(UMAP_PLOT_CLASS_ORDER) - {"Unmapped"}


def _umap_plot_class(label: object) -> str:
    key = clean(label).upper()
    if not key or key == "UNMAPPED":
        return "Unmapped"
    if key in {"EXC", "EXCITATORY", "GLUTAMATERGIC"} or "EXCIT" in key or "GLUT" in key:
        return "EXC"
    if key in {"INH", "INHIBITORY", "GABAERGIC"} or "INHIB" in key or "GABA" in key:
        return "INH"
    if key in {"OLIGO", "OPC", "MICRO", "AST", "ENDO", "VLMC"}:
        return key
    if key in {
        "ET",
        "IT",
        "NP",
        "L2-L3 IT LINC00507",
        "L3-L4 IT RORB",
        "L4-L5 IT RORB",
        "L6 CAR3",
        "L6 CT",
        "L6 IT",
        "L6B",
    } or key.startswith("L2") or key.startswith("L3") or key.startswith("L4") or key.startswith("L5") or key.startswith("L6"):
        return "EXC"
    if key in {"CHANDELIER", "LAMP5", "NDNF", "PAX6", "PVALB", "SST", "VIP"}:
        return "INH"
    return clean(label)


def _umap_crosswalk_broad_class(value: object) -> str:
    exact = {
        "excitatory neuron": "EXC",
        "inhibitory neuron": "INH",
        "astrocyte": "AST",
        "oligodendrocyte": "OLIGO",
        "opc": "OPC",
        "microglia": "MICRO",
        "endothelial": "ENDO",
        "vlmc": "VLMC",
    }
    return exact.get(clean(value).casefold(), "")


def _add_umap_broad_mapping(
    target: dict[tuple[str, str], str],
    key: tuple[str, str],
    value: str,
) -> None:
    if not value:
        return
    previous = target.get(key)
    if previous == "__AMBIGUOUS__":
        return
    if previous is not None and previous != value:
        target[key] = "__AMBIGUOUS__"
    elif previous is None:
        target[key] = value


def _load_umap_broad_fallback_map(study: str) -> dict[tuple[str, str], str]:
    """Load only explicit author broad classes from the fixed study crosswalk."""
    target: dict[tuple[str, str], str] = {}
    if study == "Allen":
        path = INPUT / "Allen/backup/label_mapping.tsv"
        table = read_mapping_table(path, "side\t")
        if not table.empty and "side" in table.columns:
            table = table[table["side"].astype(str).eq("author")]
        for record in table.to_dict("records"):
            author = clean(record.get("source_label", ""))
            broad = _umap_plot_class(record.get("common_label", ""))
            if author and broad in UMAP_ALLOWED_PLOT_CLASSES:
                _add_umap_broad_mapping(target, ("", author.casefold()), broad)
        return target
    if study == "SEAAD":
        for region_key in ("DLPFC", "MTG"):
            path = INPUT / f"SEAAD/backup/SEAAD_{region_key}_Subclass_to_canonical22.tsv"
            table = read_mapping_table(path, "source_label\t")
            for record in table.to_dict("records"):
                author = clean(record.get("source_label", ""))
                broad = _umap_plot_class(record.get("canonical22", ""))
                if author and broad in UMAP_ALLOWED_PLOT_CLASSES:
                    _add_umap_broad_mapping(target, (region_key, author.casefold()), broad)
        return target
    path = INPUT / study / "backup/author_label_to_canonical22_crosswalk.tsv"
    table = read_mapping_table(path, "dataset\t")
    for record in table.to_dict("records"):
        author = clean(record.get("author_label", ""))
        broad = _umap_crosswalk_broad_class(record.get("broad_class", ""))
        if author and broad:
            _add_umap_broad_mapping(target, ("", author.casefold()), broad)
    return target


def _apply_umap_broad_plot_classes(coordinates: Path) -> tuple[pd.DataFrame, dict[str, int]]:
    """Update only plot_class in the retained coordinate table using fixed study maps."""
    frame = pd.read_csv(coordinates, sep="\t", dtype=str, keep_default_na=False)
    fallback_maps = {
        study: _load_umap_broad_fallback_map(study)
        for study in STUDIES
    }
    restored: dict[str, int] = {}
    plot_classes = []
    for record in frame.to_dict("records"):
        study = clean(record.get("study", ""))
        canonical = clean(record.get("common_label", ""))
        broad = _umap_plot_class(canonical)
        if broad not in UMAP_ALLOWED_PLOT_CLASSES:
            author = clean(record.get("author_label", ""))
            region = clean(record.get("region", "")).upper() if study == "SEAAD" else ""
            broad = fallback_maps.get(study, {}).get((region, author.casefold()), "")
        if broad not in UMAP_ALLOWED_PLOT_CLASSES:
            broad = "Unmapped"
        previous = clean(record.get("plot_class", ""))
        if previous != broad and broad != "Unmapped":
            restored[broad] = restored.get(broad, 0) + 1
        plot_classes.append(broad)
    frame["plot_class"] = plot_classes
    frame.to_csv(coordinates, sep="\t", index=False)
    return frame, restored


def _load_umap_retained_columns() -> list[str]:
    table = pd.read_csv(RETAIN_MAP_PATH, sep="\t", dtype=str).fillna("")
    keep = table[table["status"].astype(str).str.lower().eq("kept")].copy()
    keep["old_num"] = (
        keep["old_P"].astype(str).str.replace("P", "", regex=False).astype(int)
    )
    keep["new_num"] = (
        keep["new_P"].astype(str).str.replace("P", "", regex=False).astype(int)
    )
    keep = keep.sort_values("new_num", kind="stable")
    if keep["new_num"].tolist() != list(range(1, 55)):
        raise RuntimeError("retained program map is not the contiguous 54-program map")
    return [str(value) for value in keep["old_num"].tolist()]


def _read_h5ad_obs_field(path: Path, field: str) -> dict[str, str]:
    source = ad.read_h5ad(path, backed="r")
    try:
        if field not in source.obs.columns:
            raise RuntimeError(f"{path}: missing library field {field}")
        cells = [clean(value) for value in source.obs_names]
        values = [clean(value) for value in source.obs[field].tolist()]
        return dict(zip(cells, values))
    finally:
        source.file.close()


def _mapping_details(study: str) -> dict[str, tuple[str, str, str]]:
    """Return author-label -> (canonical, status, source) from fixed tables."""
    details: dict[str, tuple[str, str, str]] = {}

    def add_table(table: pd.DataFrame, source: str, source_col: str, target_col: str) -> None:
        if table.empty or source_col not in table.columns or target_col not in table.columns:
            return
        for row in table.to_dict("records"):
            source_label = clean(row.get(source_col, ""))
            canonical = clean(row.get(target_col, ""))
            status = clean(row.get("mapping_status", ""))
            if not source_label:
                continue
            if source_label in details and details[source_label][0] != canonical:
                details[source_label] = ("", "AMBIGUOUS_MAPPING", source)
            else:
                details[source_label] = (canonical, status, source)

    if study == "Allen":
        path = INPUT / "Allen/backup/label_mapping.tsv"
        table = read_mapping_table(path, "side\t")
        if not table.empty and "side" in table.columns:
            table = table[table["side"].astype(str).eq("author")]
        add_table(table, str(path), "source_label", "common_label")
        return details
    if study == "SEAAD":
        for region in ("DLPFC", "MTG"):
            path = INPUT / f"SEAAD/backup/SEAAD_{region}_Subclass_to_canonical22.tsv"
            table = read_mapping_table(path, "source_label\t")
            add_table(table, str(path), "source_label", "canonical22")
        return details
    path = INPUT / study / "backup/author_label_to_canonical22_crosswalk.tsv"
    table = read_mapping_table(path, "dataset\t")
    add_table(table, str(path), "author_label", "canonical22")
    return details


def _apply_umap_label_mapping(study: str, frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    detail_maps = {} if study == "SEAAD" else _mapping_details(study)
    regional_maps: dict[str, dict[str, tuple[str, str, str]]] = {}
    if study == "SEAAD":
        for region_key in ("DLPFC", "MTG"):
            path = INPUT / f"SEAAD/backup/SEAAD_{region_key}_Subclass_to_canonical22.tsv"
            table = read_mapping_table(path, "source_label\t")
            regional: dict[str, tuple[str, str, str]] = {}
            if not table.empty:
                for row in table.to_dict("records"):
                    source_label = clean(row.get("source_label", ""))
                    if source_label:
                        regional[source_label] = (
                            clean(row.get("canonical22", "")),
                            clean(row.get("mapping_status", "")),
                            str(path),
                        )
            regional_maps[region_key] = regional
    canonical = []
    statuses = []
    sources = []
    for label, region in zip(result["author_label"], result["region"]):
        key = clean(label)
        if study == "SEAAD":
            # SEA-AD has separate DLPFC and MTG maps; the table choice is part of the label semantics.
            region_key = clean(region).upper()
            record = regional_maps.get(region_key, {}).get(key)
        else:
            record = detail_maps.get(key)
        if record is None or not record[0]:
            canonical.append("Unmapped")
            statuses.append(record[1] if record and record[1] else "UNMAPPED_AUTHOR_LABEL")
            sources.append(record[2] if record else "")
        else:
            canonical.append(record[0])
            statuses.append(record[1] if record[1] else "MAPPED_FIXED_TABLE")
            sources.append(record[2])
    result["common_label"] = canonical
    result["mapping_status"] = statuses
    result["mapping_source"] = sources
    return result


def _attach_umap_library_metadata(study: str, frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    result["library"] = ""
    result["unit_kind"] = ""
    result["library_source"] = ""
    result["library_is_proxy"] = False
    result["library_source_status"] = ""

    if study == "SEAAD":
        result["library"] = result["source_unit"].map(clean)
        result["unit_kind"] = "source_unit_library"
        result["library_source"] = "cell_scores.source_unit"
    elif study in {"Grubman_2019_GSE138852", "Lake_2018_GSE97930", "Schirmer_2019_MS_UCSC"}:
        result["library"] = result["source_unit"].map(clean)
        result["unit_kind"] = "paired_library_pool" if study.startswith("Grubman") else "source_unit_library"
        result["library_source"] = "cell_scores.source_unit"
    elif study == "Tran_2021_reward_cortex_CELLxGENE":
        result["library"] = result["region"].map(clean)
        result["unit_kind"] = "donor_region_sample_proxy"
        result["library_source"] = "region_proxy"
        result["library_is_proxy"] = True
    else:
        field_specs = {
            "Allen": (ALLEN_H5AD, "Specimen ID", "sample_specimen"),
            "SingleSoma_AD_PFC_CELLxGENE": (SINGLESOMA_H5AD, "Sample.ID", "sample_id"),
            "GSE144136": (GENERIC_H5ADS["GSE144136"], "sample_id", "sample_id"),
            "GSE174367": (GENERIC_H5ADS["GSE174367"], "sample_id", "sample_id"),
            "GSE291605": (GENERIC_H5ADS["GSE291605"], "library_name", "library_name"),
        }
        path, field, kind = field_specs[study]
        lookup = _read_h5ad_obs_field(path, field)
        result["library"] = result["cell_id"].map(lookup).fillna("").map(clean)
        result["unit_kind"] = kind
        result["library_source"] = f"{path}::obs[{field}]"

    result["library_source_status"] = np.where(
        result["library"].astype(str).ne(""), "AVAILABLE", "LIBRARY_SOURCE_GAP"
    )
    return result


def _external_umap_median_profiles(study: str) -> pd.DataFrame:
    path = OUTPUT / f"{study}.cell_scores.parquet"
    columns = [
        "cell_id",
        "source_unit",
        "donor",
        "region",
        "condition",
        "author_label",
        *UMAP_PROGRAMS,
    ]
    frame = pd.read_parquet(path, columns=columns)
    for column in ["cell_id", "source_unit", "donor", "region", "condition", "author_label"]:
        frame[column] = frame[column].map(clean)
    frame = _attach_umap_library_metadata(study, frame)
    frame = _apply_umap_label_mapping(study, frame)
    keys = ["library", "donor", "author_label", "region", "condition"]
    rows = []
    for key, positions in frame.groupby(keys, sort=False, dropna=False).indices.items():
        subset = frame.iloc[positions]
        values = subset[UMAP_PROGRAMS].to_numpy(dtype=np.float64)
        median = np.median(values, axis=0)
        profile_status = "OK"
        if not np.isfinite(median).all():
            profile_status = "NONFINITE_PROFILE"
        elif np.ptp(median) == 0.0:
            profile_status = "CONSTANT_PROFILE"
        if subset["library_source_status"].iloc[0] == "LIBRARY_SOURCE_GAP":
            profile_status = "LIBRARY_SOURCE_GAP" if profile_status == "OK" else f"LIBRARY_SOURCE_GAP;{profile_status}"
        library, donor, author_label, region, condition = key
        row = {
            "point_id": "::".join([study, library, donor, author_label, region, condition]),
            "point_kind": "external_library_celltype_median",
            "study": study,
            "library": library,
            "source_unit": clean(subset["source_unit"].iloc[0]),
            "donor": donor,
            "author_label": author_label,
            "common_label": clean(subset["common_label"].iloc[0]),
            "mapping_status": clean(subset["mapping_status"].iloc[0]),
            "mapping_source": clean(subset["mapping_source"].iloc[0]),
            "region": region,
            "condition": condition,
            "n_cells": int(len(subset)),
            "unit_kind": clean(subset["unit_kind"].iloc[0]),
            "library_source": clean(subset["library_source"].iloc[0]),
            "library_is_proxy": bool(subset["library_is_proxy"].iloc[0]),
            "library_source_status": clean(subset["library_source_status"].iloc[0]),
            "profile_status": profile_status,
        }
        row.update({name: float(value) for name, value in zip(UMAP_PROGRAMS, median)})
        rows.append(row)
    return pd.DataFrame(rows, columns=UMAP_PROFILE_COLUMNS)


def _reference_umap_median_profiles() -> pd.DataFrame:
    old_columns = _load_umap_retained_columns()
    source_columns = [str(value) for value in range(1, 61)]
    path = PROJECT / "archived/results/crossregion_v1/cell_program_region_subclass.parquet"
    scores = pd.read_parquet(path, columns=source_columns)
    if "__index_level_0__" in scores.columns:
        cell_ids = [clean(value) for value in scores["__index_level_0__"].tolist()]
    else:
        cell_ids = [clean(value) for value in scores.index.tolist()]
    obs = pd.read_csv(
        DISCOVERY_OBS_PATH,
        index_col=0,
        usecols=["Unnamed: 0", "subclass"],
        dtype=str,
    ).fillna("")
    obs.index = [clean(value) for value in obs.index]
    subclasses = obs.reindex(cell_ids)["subclass"].map(clean).to_numpy(dtype=str)
    values = scores[source_columns].to_numpy(dtype=np.float64)
    row_sum = values.sum(axis=1, keepdims=True)
    normalized = np.divide(values, row_sum, out=np.zeros_like(values), where=row_sum > 0.0)
    selected = normalized[:, np.asarray(old_columns, dtype=np.int64) - 1]
    rows = []
    for subclass in sorted(set(subclasses)):
        if not subclass:
            continue
        median = np.median(selected[subclasses == subclass], axis=0)
        profile_status = "OK" if np.isfinite(median).all() else "NONFINITE_PROFILE"
        if profile_status == "OK" and np.ptp(median) == 0.0:
            profile_status = "CONSTANT_PROFILE"
        row = {
            "point_id": f"reference_1M::{subclass}",
            "point_kind": "reference_subclass_median",
            "study": "reference_1M",
            "library": "",
            "source_unit": "",
            "donor": "",
            "author_label": subclass,
            "common_label": subclass,
            "mapping_status": "REFERENCE_CANONICAL22",
            "mapping_source": str(DISCOVERY_OBS_PATH),
            "region": "",
            "condition": "",
            "n_cells": int(np.sum(subclasses == subclass)),
            "unit_kind": "reference_subclass_prototype",
            "library_source": str(DISCOVERY_OBS_PATH),
            "library_is_proxy": False,
            "library_source_status": "REFERENCE_NOT_LIBRARY_GROUPED",
            "profile_status": profile_status,
        }
        row.update({name: float(value) for name, value in zip(UMAP_PROGRAMS, median)})
        rows.append(row)
    return pd.DataFrame(rows, columns=UMAP_PROFILE_COLUMNS)


def _render_umap_coordinates(profiles: pd.DataFrame) -> tuple[int, int, int, int]:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D

    profiles = profiles.copy()
    for column in ["point_kind", "study", "common_label"]:
        profiles[column] = profiles[column].map(clean)
    if "plot_class" not in profiles.columns:
        profiles["plot_class"] = profiles["common_label"].map(_umap_plot_class)
    else:
        profiles["plot_class"] = profiles["plot_class"].map(_umap_plot_class)
    x = pd.to_numeric(profiles["UMAP1"], errors="coerce").to_numpy(dtype=float)
    y = pd.to_numeric(profiles["UMAP2"], errors="coerce").to_numpy(dtype=float)
    if not np.isfinite(x).all() or not np.isfinite(y).all():
        raise RuntimeError("UMAP coordinate table contains non-finite coordinates")

    reference_mask = profiles["point_kind"].eq("reference_subclass_median").to_numpy()
    hidden_unmapped = (~reference_mask) & profiles["plot_class"].eq("Unmapped").to_numpy()
    display_mask = ~hidden_unmapped
    external_display = (~reference_mask) & display_mask
    retained_points = int(len(profiles))
    displayed_points = int(display_mask.sum())
    hidden_points = int(hidden_unmapped.sum())
    reference_points = int(reference_mask.sum())
    visible_x = x[display_mask]
    visible_y = y[display_mask]
    x_span = float(np.ptp(visible_x))
    y_span = float(np.ptp(visible_y))
    x_pad = 0.05 * (x_span if x_span > 0 else 1.0)
    y_pad = 0.05 * (y_span if y_span > 0 else 1.0)
    limits = (
        visible_x.min() - x_pad,
        visible_x.max() + x_pad,
        visible_y.min() - y_pad,
        visible_y.max() + y_pad,
    )

    figure, axis = plt.subplots(figsize=(7.2, 6.4))
    figure.patch.set_facecolor("white")
    axis.set_facecolor("white")
    present_classes = sorted(profiles.loc[external_display, "plot_class"].astype(str).unique())
    plot_classes = [value for value in UMAP_PLOT_CLASS_ORDER if value in present_classes]
    plot_classes.extend(value for value in present_classes if value not in plot_classes)
    extra_colors = plt.cm.tab20(np.linspace(0, 1, max(len(plot_classes), 1)))
    plot_colors = {
        value: UMAP_PLOT_CLASS_COLORS.get(value, extra_colors[index])
        for index, value in enumerate(plot_classes)
    }
    for plot_class in plot_classes:
        mask = external_display & profiles["plot_class"].eq(plot_class).to_numpy()
        if mask.any():
            axis.scatter(
                x[mask],
                y[mask],
                s=10,
                alpha=0.58,
                color=plot_colors[plot_class],
                edgecolors="none",
                label=plot_class,
                rasterized=True,
            )
    axis.scatter(
        x[reference_mask],
        y[reference_mask],
        s=62,
        marker="*",
        facecolors="white",
        edgecolors="black",
        linewidths=0.9,
        zorder=4,
    )
    reference_label_offsets = {
        "L6 CAR3": (10, 14),
        "PAX6": (8, -14),
        "LAMP5": (-24, -12),
        "ENDO": (-26, -14),
        "MICRO": (8, -12),
        "NDNF": (8, -10),
        "ET": (-18, 12),
        "VLMC": (-22, 12),
        "OLIGO": (8, 10),
        "SST": (8, -12),
        "L4-L5 IT RORB": (-50, 10),
        "VIP": (-16, 12),
        "NP": (-18, 12),
        "AST": (-10, 12),
        "L2-L3 IT LINC00507": (-74, 12),
        "L6 IT": (-20, 12),
        "PVALB": (10, -14),
        "OPC": (-12, 18),
        "CHANDELIER": (10, 24),
        "L3-L4 IT RORB": (10, 24),
        "L6 CT": (-22, 12),
        "L6B": (8, 10),
    }
    for reference_index in np.flatnonzero(reference_mask):
        label = clean(profiles.iloc[reference_index].get("common_label", ""))
        if not label:
            label = clean(profiles.iloc[reference_index].get("author_label", ""))
        if not label:
            continue
        offset_x, offset_y = reference_label_offsets.get(label, (8, 8))
        axis.annotate(
            label,
            xy=(x[reference_index], y[reference_index]),
            xytext=(offset_x, offset_y),
            textcoords="offset points",
            fontsize=6.5,
            color="#28323d",
            ha="left" if offset_x >= 0 else "right",
            va="bottom" if offset_y >= 0 else "top",
            arrowprops={
                "arrowstyle": "-",
                "lw": 0.45,
                "color": "#7b8794",
                "shrinkA": 2,
                "shrinkB": 2,
            },
            zorder=5,
        )
    axis.set_xlim(limits[0], limits[1])
    axis.set_ylim(limits[2], limits[3])
    axis.set_box_aspect(1)
    axis.set_title("External 54-GEP profiles", fontsize=12, pad=10)
    axis.grid(False)
    for spine in axis.spines.values():
        spine.set_visible(False)
    axis.tick_params(axis="both", which="both", length=0, labelsize=8, colors="#56616f")
    axis.set_xticks([])
    axis.set_yticks([])
    axis.annotate(
        "",
        xy=(0.16, 0.08),
        xytext=(0.05, 0.08),
        xycoords="axes fraction",
        arrowprops={"arrowstyle": "->", "lw": 1.1, "color": "#56616f"},
    )
    axis.annotate(
        "",
        xy=(0.05, 0.19),
        xytext=(0.05, 0.08),
        xycoords="axes fraction",
        arrowprops={"arrowstyle": "->", "lw": 1.1, "color": "#56616f"},
    )
    axis.text(0.17, 0.045, "UMAP1", transform=axis.transAxes, fontsize=8, color="#56616f", ha="left")
    axis.text(0.025, 0.20, "UMAP2", transform=axis.transAxes, fontsize=8, color="#56616f", rotation=90, va="bottom")
    legend_handles = [
        Line2D([], [], marker="o", linestyle="None", markerfacecolor=plot_colors[value], markeredgecolor="none", markersize=5, label=value)
        for value in plot_classes
    ]
    legend_handles.append(
        Line2D([], [], marker="*", linestyle="None", markerfacecolor="white", markeredgecolor="black", markersize=8, label="1M reference subclass")
    )
    axis.legend(
        handles=legend_handles,
        title="Broad subclass",
        title_fontsize=9,
        fontsize=8,
        ncol=1,
        loc="upper left",
        bbox_to_anchor=(1.02, 1.0),
        frameon=False,
        borderaxespad=0.0,
        handletextpad=0.5,
        labelspacing=0.45,
    )
    figure.subplots_adjust(left=0.08, right=0.76, bottom=0.08, top=0.90)
    pdf_path = OUTPUT / "library_celltype_umap_dual_view.pdf"
    png_path = OUTPUT / "library_celltype_umap_dual_view.png"
    figure.savefig(pdf_path, bbox_inches="tight", facecolor="white")
    figure.savefig(png_path, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(figure)
    return displayed_points, hidden_points, retained_points, reference_points


def run_library_celltype_umap() -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    import umap

    external_profiles = []
    point_counts = {}
    cell_counts = {}
    for study in STUDIES:
        profile = _external_umap_median_profiles(study)
        external_profiles.append(profile)
        point_counts[study] = int(len(profile))
        cell_counts[study] = int(profile["n_cells"].sum()) if len(profile) else 0
    external = pd.concat(external_profiles, ignore_index=True)
    reference = _reference_umap_median_profiles()
    profiles = pd.concat([reference, external], ignore_index=True)
    profiles["plot_class"] = profiles["common_label"].map(_umap_plot_class)
    matrix = profiles[UMAP_PROGRAMS].to_numpy(dtype=np.float64)
    if not np.isfinite(matrix).all():
        raise RuntimeError("UMAP input contains non-finite median profile values")

    reducer = umap.UMAP(
        n_neighbors=15,
        min_dist=0.1,
        n_components=2,
        metric="euclidean",
        random_state=42,
    )
    coordinates = reducer.fit_transform(matrix)
    profiles["UMAP1"] = coordinates[:, 0]
    profiles["UMAP2"] = coordinates[:, 1]

    profile_path = OUTPUT / "library_celltype_median_profiles.tsv"
    reference_path = OUTPUT / "reference_subclass_median_profiles.tsv"
    coordinate_path = OUTPUT / "library_celltype_umap_coordinates.tsv"
    pdf_path = OUTPUT / "library_celltype_umap_dual_view.pdf"
    png_path = OUTPUT / "library_celltype_umap_dual_view.png"
    external.to_csv(profile_path, sep="\t", index=False)
    reference.to_csv(reference_path, sep="\t", index=False)
    profiles.to_csv(coordinate_path, sep="\t", index=False)

    _render_umap_coordinates(profiles)

    print("UMAP_COMPLETE")
    print(f"reference_points\t{len(reference)}\t{int(reference['n_cells'].sum())}")
    for study in STUDIES:
        print(f"study_points\t{study}\t{point_counts[study]}\t{cell_counts[study]}")
    print(f"output\t{profile_path}")
    print(f"output\t{reference_path}")
    print(f"output\t{coordinate_path}")
    print(f"output\t{pdf_path}")
    print(f"output\t{png_path}")


MEDIAN_CORRELATION_COLUMNS = [
    *UMAP_PROFILE_COLUMNS[:18],
    "corresponding_reference_subclass",
    "corresponding_rho",
    "corresponding_rho_reason",
    "highest_other_reference_subclass",
    "highest_other_rho",
    "highest_other_rho_reason",
    "correlation_status",
]


def _umap_profile_reason(values: np.ndarray, prefix: str) -> str:
    array = np.asarray(values, dtype=np.float64)
    if not np.isfinite(array).all():
        return f"{prefix}_NONFINITE_PROFILE"
    if array.size < 2:
        return f"{prefix}_INSUFFICIENT_PROGRAMS"
    if np.ptp(array) == 0.0:
        return f"{prefix}_CONSTANT_PROFILE"
    return ""


def _umap_spearman(a: np.ndarray, b: np.ndarray) -> float:
    x = np.asarray(a, dtype=np.float64)
    y = np.asarray(b, dtype=np.float64)
    if x.size != y.size or x.size < 2:
        return float("nan")
    if not np.isfinite(x).all() or not np.isfinite(y).all():
        return float("nan")
    rx = pd.Series(x).rank(method="average").to_numpy(dtype=np.float64)
    ry = pd.Series(y).rank(method="average").to_numpy(dtype=np.float64)
    rx -= rx.mean()
    ry -= ry.mean()
    denominator = float(np.sqrt(np.square(rx).sum() * np.square(ry).sum()))
    if denominator == 0.0:
        return float("nan")
    return float(np.dot(rx, ry) / denominator)


def _clean_median_profile_frame(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    metadata_columns = UMAP_PROFILE_COLUMNS[:18]
    for column in metadata_columns:
        if column not in result.columns:
            result[column] = ""
    for column in metadata_columns:
        if column == "n_cells":
            result[column] = pd.to_numeric(result[column], errors="coerce").fillna(0).astype(np.int64)
        elif column == "library_is_proxy":
            result[column] = result[column].map(
                lambda value: clean(value).lower() in {"true", "1", "yes"}
            )
        else:
            result[column] = result[column].map(clean)
    for column in UMAP_PROGRAMS:
        result[column] = pd.to_numeric(result[column], errors="coerce")
    return result


def run_library_celltype_median_correlations() -> None:
    profile_path = OUTPUT / "library_celltype_median_profiles.tsv"
    reference_path = OUTPUT / "reference_subclass_median_profiles.tsv"
    correlation_path = OUTPUT / "library_celltype_median_correlations.tsv"
    external = _clean_median_profile_frame(pd.read_csv(profile_path, sep="\t"))
    reference = _clean_median_profile_frame(pd.read_csv(reference_path, sep="\t"))

    reference_values: dict[str, np.ndarray] = {}
    reference_reasons: dict[str, str] = {}
    for record in reference.to_dict("records"):
        label = clean(record.get("common_label", ""))
        if not label or label in reference_values:
            continue
        values = np.asarray([record.get(name, np.nan) for name in UMAP_PROGRAMS], dtype=np.float64)
        reference_values[label] = values
        reference_reasons[label] = _umap_profile_reason(values, "REFERENCE")

    rows = []
    for record in external.to_dict("records"):
        row = {column: record.get(column, "") for column in UMAP_PROFILE_COLUMNS[:18]}
        row["point_id"] = clean(row.get("point_id", ""))
        row["point_kind"] = clean(row.get("point_kind", ""))
        row["study"] = clean(row.get("study", ""))
        row["library"] = clean(row.get("library", ""))
        row["source_unit"] = clean(row.get("source_unit", ""))
        row["donor"] = clean(row.get("donor", ""))
        row["author_label"] = clean(row.get("author_label", ""))
        row["common_label"] = clean(row.get("common_label", ""))
        row["mapping_status"] = clean(row.get("mapping_status", ""))
        row["mapping_source"] = clean(row.get("mapping_source", ""))
        row["region"] = clean(row.get("region", ""))
        row["condition"] = clean(row.get("condition", ""))
        n_cells_value = pd.to_numeric(row.get("n_cells", 0), errors="coerce")
        row["n_cells"] = int(n_cells_value) if pd.notna(n_cells_value) else 0
        row["unit_kind"] = clean(row.get("unit_kind", ""))
        row["library_source"] = clean(row.get("library_source", ""))
        row["library_is_proxy"] = clean(row.get("library_is_proxy", "")).lower() in {"true", "1", "yes"}
        row["library_source_status"] = clean(row.get("library_source_status", ""))
        row["profile_status"] = clean(row.get("profile_status", ""))
        values = np.asarray([record.get(name, np.nan) for name in UMAP_PROGRAMS], dtype=np.float64)
        external_reason = _umap_profile_reason(values, "EXTERNAL")
        common_label = row["common_label"]
        row.update(
            {
                "corresponding_reference_subclass": "",
                "corresponding_rho": np.nan,
                "corresponding_rho_reason": "",
                "highest_other_reference_subclass": "",
                "highest_other_rho": np.nan,
                "highest_other_rho_reason": "",
                "correlation_status": "",
            }
        )
        if external_reason:
            row["corresponding_rho_reason"] = external_reason
            row["highest_other_rho_reason"] = external_reason
            row["correlation_status"] = "UNDEFINED_EXTERNAL_PROFILE"
        elif not common_label or common_label == "Unmapped" or common_label not in reference_values:
            reason = "UNMAPPED_AUTHOR_LABEL" if not common_label or common_label == "Unmapped" else "NO_CORRESPONDING_REFERENCE_SUBCLASS"
            row["corresponding_rho_reason"] = reason
            row["highest_other_rho_reason"] = reason
            row["correlation_status"] = "UNDEFINED_NO_CORRESPONDING_REFERENCE"
        else:
            row["corresponding_reference_subclass"] = common_label
            reference_reason = reference_reasons[common_label]
            if reference_reason:
                row["corresponding_rho_reason"] = reference_reason
            else:
                corresponding_rho = _umap_spearman(values, reference_values[common_label])
                if np.isfinite(corresponding_rho):
                    row["corresponding_rho"] = corresponding_rho
                    row["corresponding_rho_reason"] = "DEFINED"
                else:
                    row["corresponding_rho_reason"] = "UNDEFINED_CORRESPONDING_RHO"

            other_rhos = []
            for label, reference_vector in reference_values.items():
                if label == common_label or reference_reasons[label]:
                    continue
                other_rho = _umap_spearman(values, reference_vector)
                if np.isfinite(other_rho):
                    other_rhos.append((label, other_rho))
            if other_rhos:
                other_label, other_rho = max(other_rhos, key=lambda item: item[1])
                row["highest_other_reference_subclass"] = other_label
                row["highest_other_rho"] = other_rho
                row["highest_other_rho_reason"] = "DEFINED"
            else:
                row["highest_other_rho_reason"] = "NO_DEFINED_OTHER_REFERENCE_RHO"
            if np.isfinite(row["corresponding_rho"]):
                row["correlation_status"] = "DEFINED" if np.isfinite(row["highest_other_rho"]) else "PARTIAL_DEFINED"
            else:
                row["correlation_status"] = "UNDEFINED_CORRESPONDING_RHO"
        rows.append(row)

    correlations = pd.DataFrame(rows, columns=MEDIAN_CORRELATION_COLUMNS)
    correlations.to_csv(correlation_path, sep="\t", index=False, na_rep="NA")
    print("MEDIAN_CORRELATIONS_COMPLETE")
    print(f"rho_rows\t{len(correlations)}")
    print(f"defined_corresponding_rho\t{int(correlations['corresponding_rho'].notna().sum())}")
    print(f"undefined_or_unmapped_rows\t{int(correlations['corresponding_rho'].isna().sum())}")
    for column, label in (
        ("corresponding_rho_reason", "corresponding_reason"),
        ("highest_other_rho_reason", "highest_other_reason"),
    ):
        for reason, count in correlations[column].value_counts(dropna=False).items():
            reason = clean(reason)
            if reason:
                print(f"{label}\t{reason}\t{int(count)}")
    print(f"output\t{correlation_path}")


def render_library_celltype_umap() -> None:
    coordinate_path = OUTPUT / "library_celltype_umap_coordinates.tsv"
    profiles = pd.read_csv(coordinate_path, sep="\t")
    displayed_points, hidden_points, retained_points, reference_points = _render_umap_coordinates(profiles)
    print("UMAP_RENDER_COMPLETE")
    print(f"displayed_points\t{displayed_points}")
    print(f"hidden_unmapped_points\t{hidden_points}")
    print(f"retained_points\t{retained_points}")
    print(f"reference_points\t{reference_points}")
    print(f"output\t{OUTPUT / 'library_celltype_umap_dual_view.pdf'}")
    print(f"output\t{OUTPUT / 'library_celltype_umap_dual_view.png'}")


HEATMAP_STUDY_LABELS = {
    "Allen": "Allen",
    "SEAAD": "SEA-AD",
    "Grubman_2019_GSE138852": "Grubman",
    "Lake_2018_GSE97930": "Lake",
    "Schirmer_2019_MS_UCSC": "Schirmer",
    "SingleSoma_AD_PFC_CELLxGENE": "SingleSoma",
    "Tran_2021_reward_cortex_CELLxGENE": "Tran",
    "GSE144136": "Nagy",
    "GSE174367": "Morabito",
    "GSE291605": "Mural",
}
HEATMAP_COLUMNS = [
    "study",
    "study_display",
    "study_order",
    "author_label",
    "author_label_display",
    "original_author_order",
    "author_type_order",
    "column_key",
    "column_order",
    "common_label",
    "mapping_status",
    "reference_subclass",
    "reference_order",
    "spearman_rho",
    "contributing_group_count",
    "valid_group_count",
    "source_unit_kind",
    "aggregation_unit_kind",
    "aggregation_unit_count",
    "valid_aggregation_unit_count",
    "n_cells_total",
    "n_cells_valid",
    "region_coverage",
    "condition_coverage",
    "na_reason",
    "sorting_peak_reference",
    "sorting_peak_rho",
]


def _heatmap_join_values(values: Iterable[object]) -> str:
    cleaned = sorted({clean(value) for value in values})
    cleaned = [value if value else "UNSPECIFIED" for value in cleaned]
    return ";".join(cleaned)


def _heatmap_reason_text(reasons: Iterable[object]) -> str:
    counts: dict[str, int] = {}
    for value in reasons:
        reason = clean(value)
        if reason:
            counts[reason] = counts.get(reason, 0) + 1
    return ";".join(
        f"{reason}({counts[reason]})" for reason in sorted(counts)
    )


def _heatmap_aggregation_unit(record: dict[str, object]) -> tuple[str, str]:
    unit_kind = clean(record.get("unit_kind", ""))
    library = clean(record.get("library", ""))
    donor = clean(record.get("donor", ""))
    region = clean(record.get("region", ""))
    point_id = clean(record.get("point_id", ""))
    if unit_kind == "paired_library_pool":
        return f"pool::{library or record.get('source_unit', '')}", "paired_library_pool"
    if unit_kind == "donor_region_sample_proxy":
        if donor:
            return f"donor::{donor}", "donor_with_region_sample_proxy_libraries"
        return f"donor_region::{donor}::{region}", "donor_region_sample_proxy"
    if donor:
        return f"donor::{donor}", "donor"
    if library:
        return f"sample::{library}", unit_kind or "sample"
    return f"point::{point_id}", unit_kind or "point"


def _heatmap_column_definitions(external: pd.DataFrame) -> list[dict[str, object]]:
    first_seen: dict[tuple[str, str], int] = {}
    for index, record in enumerate(external.to_dict("records")):
        key = (clean(record.get("study", "")), clean(record.get("author_label", "")))
        if key not in first_seen:
            first_seen[key] = index
    definitions = []
    for study_order, study in enumerate(STUDIES, start=1):
        keys = [key for key in first_seen if key[0] == study]
        keys.sort(key=lambda key: first_seen[key])
        for author_type_order, key in enumerate(keys, start=1):
            definitions.append(
                {
                    "study": study,
                    "study_display": HEATMAP_STUDY_LABELS.get(study, study),
                    "study_order": study_order,
                    "author_label": key[1],
                    "author_label_display": key[1] or "Missing original annotation",
                    "original_author_order": author_type_order,
                    "author_type_order": author_type_order,
                    "column_key": f"{study}::{key[1]}",
                    "column_order": len(definitions) + 1,
                }
            )
    return definitions


def _order_heatmap_columns(
    table: pd.DataFrame, definitions: list[dict[str, object]]
) -> tuple[pd.DataFrame, list[dict[str, object]]]:
    peak_info: dict[str, tuple[int, float, str]] = {}
    for definition in definitions:
        key = definition["column_key"]
        subset = table[table["column_key"].eq(key)]
        values = pd.to_numeric(subset["spearman_rho"], errors="coerce").to_numpy(dtype=float)
        rows = pd.to_numeric(subset["reference_order"], errors="coerce").to_numpy(dtype=int)
        finite = np.isfinite(values)
        if finite.any():
            peak_rho = float(np.max(values[finite]))
            peak_rows = rows[finite & (values == peak_rho)]
            peak_order = int(np.min(peak_rows))
            peak_reference = clean(
                subset.loc[subset["reference_order"].eq(peak_order), "reference_subclass"].iloc[0]
            )
            peak_info[key] = (peak_order, peak_rho, peak_reference)
        else:
            peak_info[key] = (10**9, float("nan"), "")

    ordered = []
    for study in STUDIES:
        study_definitions = [value for value in definitions if value["study"] == study]
        study_definitions.sort(
            key=lambda value: (
                peak_info[value["column_key"]][0],
                -peak_info[value["column_key"]][1]
                if np.isfinite(peak_info[value["column_key"]][1])
                else 0.0,
                value["original_author_order"],
            )
        )
        for author_type_order, definition in enumerate(study_definitions, start=1):
            updated = dict(definition)
            updated["author_type_order"] = author_type_order
            updated["column_order"] = len(ordered) + 1
            peak_order, peak_rho, peak_reference = peak_info[definition["column_key"]]
            updated["sorting_peak_reference"] = peak_reference
            updated["sorting_peak_rho"] = peak_rho
            ordered.append(updated)

    by_key = {value["column_key"]: value for value in ordered}
    for column in [
        "study_display",
        "study_order",
        "author_label_display",
        "original_author_order",
        "author_type_order",
        "column_order",
        "sorting_peak_reference",
        "sorting_peak_rho",
    ]:
        table[column] = table["column_key"].map(
            {key: value.get(column, "") for key, value in by_key.items()}
        )
    table = table.sort_values(
        ["reference_order", "column_order"], kind="stable"
    ).reset_index(drop=True)
    return table, ordered


def _heatmap_table_from_medians(
    external: pd.DataFrame,
    reference: pd.DataFrame,
) -> tuple[pd.DataFrame, list[dict[str, object]]]:
    external = _clean_median_profile_frame(external)
    reference = _clean_median_profile_frame(reference)
    definitions = _heatmap_column_definitions(external)
    definition_by_key = {value["column_key"]: value for value in definitions}
    reference_labels = []
    reference_values: dict[str, np.ndarray] = {}
    reference_reasons: dict[str, str] = {}
    for record in reference.to_dict("records"):
        label = clean(record.get("common_label", ""))
        if not label or label in reference_values:
            continue
        values = np.asarray(
            [record.get(name, np.nan) for name in UMAP_PROGRAMS], dtype=np.float64
        )
        reference_labels.append(label)
        reference_values[label] = values
        reference_reasons[label] = _umap_profile_reason(values, "REFERENCE")

    grouped: dict[tuple[str, str], dict[str, dict[str, object]]] = {}
    for record in external.to_dict("records"):
        study = clean(record.get("study", ""))
        author_label = clean(record.get("author_label", ""))
        column_key = f"{study}::{author_label}"
        if column_key not in definition_by_key:
            continue
        unit_key, unit_kind = _heatmap_aggregation_unit(record)
        group_key = (study, author_label)
        grouped.setdefault(group_key, {}).setdefault(
            unit_key,
            {
                "unit_kind": unit_kind,
                "records": [],
            },
        )["records"].append(record)

    rows = []
    for definition in definitions:
        group_key = (definition["study"], definition["author_label"])
        units = grouped.get(group_key, {})
        for reference_order, reference_label in enumerate(reference_labels, start=1):
            unit_medians = []
            unit_reasons = []
            contributing_group_count = 0
            valid_group_count = 0
            n_cells_total = 0
            n_cells_valid = 0
            regions = []
            conditions = []
            common_labels = []
            mapping_statuses = []
            source_unit_kinds = []
            valid_unit_count = 0
            for unit in units.values():
                records = unit["records"]
                unit_rhos = []
                unit_reasons_for_reference = []
                for record in records:
                    contributing_group_count += 1
                    n_cells = int(record.get("n_cells", 0))
                    n_cells_total += n_cells
                    regions.append(record.get("region", ""))
                    conditions.append(record.get("condition", ""))
                    common_labels.append(record.get("common_label", ""))
                    mapping_statuses.append(record.get("mapping_status", ""))
                    source_unit_kinds.append(record.get("unit_kind", ""))
                    values = np.asarray(
                        [record.get(name, np.nan) for name in UMAP_PROGRAMS],
                        dtype=np.float64,
                    )
                    external_reason = _umap_profile_reason(values, "EXTERNAL")
                    if external_reason:
                        rho = float("nan")
                        reason = external_reason
                    else:
                        reference_reason = reference_reasons[reference_label]
                        if reference_reason:
                            rho = float("nan")
                            reason = reference_reason
                        else:
                            rho = _umap_spearman(values, reference_values[reference_label])
                            reason = "" if np.isfinite(rho) else "UNDEFINED_CORRELATION"
                    if np.isfinite(rho):
                        unit_rhos.append(rho)
                        valid_group_count += 1
                        n_cells_valid += n_cells
                    else:
                        unit_reasons_for_reference.append(reason)
                        unit_reasons.append(reason)
                if unit_rhos:
                    unit_medians.append(float(np.median(np.asarray(unit_rhos, dtype=float))))
                    valid_unit_count += 1
                else:
                    unit_medians.append(float("nan"))
            valid_unit_medians = [value for value in unit_medians if np.isfinite(value)]
            rho = float(np.median(np.asarray(valid_unit_medians, dtype=float))) if valid_unit_medians else float("nan")
            na_reason = "" if np.isfinite(rho) else (
                _heatmap_reason_text(unit_reasons) if unit_reasons else "NO_CONTRIBUTING_GROUPS"
            )
            rows.append(
                {
                    **definition,
                    "common_label": _heatmap_join_values(common_labels),
                    "mapping_status": _heatmap_join_values(mapping_statuses),
                    "reference_subclass": reference_label,
                    "reference_order": reference_order,
                    "spearman_rho": rho,
                    "contributing_group_count": contributing_group_count,
                    "valid_group_count": valid_group_count,
                    "source_unit_kind": _heatmap_join_values(source_unit_kinds),
                    "aggregation_unit_kind": _heatmap_join_values(
                        unit.get("unit_kind", "") for unit in units.values()
                    ),
                    "aggregation_unit_count": len(units),
                    "valid_aggregation_unit_count": valid_unit_count,
                    "n_cells_total": n_cells_total,
                    "n_cells_valid": n_cells_valid,
                    "region_coverage": _heatmap_join_values(regions),
                    "condition_coverage": _heatmap_join_values(conditions),
                    "na_reason": na_reason,
                }
            )
    table = pd.DataFrame(rows, columns=HEATMAP_COLUMNS)
    return _order_heatmap_columns(table, definitions)


def _render_study_celltype_heatmap(
    table: pd.DataFrame, definitions: list[dict[str, object]]
) -> tuple[int, int]:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Patch

    reference_order = sorted(table["reference_order"].unique())
    reference_labels = (
        table[["reference_order", "reference_subclass"]]
        .drop_duplicates()
        .sort_values("reference_order")["reference_subclass"]
        .tolist()
    )
    column_order = [value["column_key"] for value in definitions]
    column_index = {value: index for index, value in enumerate(column_order)}
    row_index = {value: index for index, value in enumerate(reference_order)}
    matrix = np.full((len(reference_order), len(column_order)), np.nan, dtype=float)
    for record in table.to_dict("records"):
        rho = pd.to_numeric(record.get("spearman_rho"), errors="coerce")
        if pd.notna(rho):
            matrix[row_index[record["reference_order"]], column_index[record["column_key"]]] = float(rho)
    masked = np.ma.masked_invalid(matrix)
    cmap = plt.get_cmap("RdBu_r").copy()
    cmap.set_bad("#eeeeee")
    figure_width = max(36.0, 0.24 * len(column_order) + 8.0)
    figure, axis = plt.subplots(figsize=(figure_width, 10.5))
    mesh = axis.pcolormesh(
        masked,
        cmap=cmap,
        vmin=-1.0,
        vmax=1.0,
        shading="flat",
        edgecolors="white",
        linewidth=0.12,
    )
    axis.set_xlim(0, len(column_order))
    axis.set_ylim(len(reference_order), 0)
    axis.set_xticks(np.arange(len(column_order)) + 0.5)
    axis.set_xticklabels([value["author_label_display"] for value in definitions])
    axis.set_yticks(np.arange(len(reference_order)) + 0.5)
    axis.set_yticklabels(reference_labels)
    axis.tick_params(axis="x", labelrotation=90, labelsize=6.5, length=0, pad=2)
    axis.tick_params(axis="y", labelsize=8, length=0, pad=3)
    axis.set_xlabel("Original author cell type (study-scoped)", fontsize=9, labelpad=10)
    axis.set_ylabel("1M reference subclass", fontsize=10, labelpad=8)
    axis.set_title(
        "Study × author cell type | 1M reference rho\n"
        "Columns ordered by peak reference row within study",
        fontsize=12,
        pad=34,
    )
    for spine in axis.spines.values():
        spine.set_visible(False)
    start = 0
    for study in STUDIES:
        count = sum(value["study"] == study for value in definitions)
        if count == 0:
            continue
        axis.axvline(start, color="#4b5563", linewidth=0.8)
        axis.text(
            start + count / 2,
            1.015,
            HEATMAP_STUDY_LABELS.get(study, study),
            transform=axis.get_xaxis_transform(),
            ha="center",
            va="bottom",
            fontsize=8,
            color="#26323e",
        )
        start += count
    axis.axvline(len(column_order), color="#4b5563", linewidth=0.8)
    colorbar = figure.colorbar(mesh, ax=axis, fraction=0.012, pad=0.015)
    colorbar.set_label("Spearman rho", fontsize=9)
    colorbar.set_ticks([-1, -0.5, 0, 0.5, 1])
    axis.legend(
        handles=[Patch(facecolor="#eeeeee", edgecolor="#c7cbd1", label="NA")],
        loc="upper left",
        bbox_to_anchor=(1.005, 1.0),
        frameon=False,
        fontsize=8,
        borderaxespad=0.0,
    )
    figure.subplots_adjust(left=0.12, right=0.965, bottom=0.31, top=0.86)
    pdf_path = OUTPUT / "study_celltype_reference_correlation_heatmap.pdf"
    png_path = OUTPUT / "study_celltype_reference_correlation_heatmap.png"
    figure.savefig(pdf_path, bbox_inches="tight", facecolor="white")
    figure.savefig(png_path, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(figure)
    defined_cells = int(np.isfinite(matrix).sum())
    return defined_cells, int(matrix.size - defined_cells)


def run_study_celltype_reference_heatmap() -> None:
    profile_path = OUTPUT / "library_celltype_median_profiles.tsv"
    reference_path = OUTPUT / "reference_subclass_median_profiles.tsv"
    heatmap_path = OUTPUT / "study_celltype_reference_correlation_heatmap.tsv"
    external = pd.read_csv(profile_path, sep="\t")
    reference = pd.read_csv(reference_path, sep="\t")
    table, definitions = _heatmap_table_from_medians(external, reference)
    table.to_csv(heatmap_path, sep="\t", index=False, na_rep="NA")
    defined_cells, na_cells = _render_study_celltype_heatmap(table, definitions)
    print("HEATMAP_COMPLETE")
    print(f"rows\t{len(table)}")
    print(f"columns\t{len(definitions)}")
    print(f"cells\t{len(table)}")
    print(f"defined_cells\t{defined_cells}")
    print(f"na_cells\t{na_cells}")
    print(f"external_groups\t{len(external)}")
    for reason, count in table.loc[table["spearman_rho"].isna(), "na_reason"].value_counts().items():
        if clean(reason):
            print(f"na_reason\t{reason}\t{int(count)}")
    print(f"output\t{heatmap_path}")
    print(f"output\t{OUTPUT / 'study_celltype_reference_correlation_heatmap.pdf'}")
    print(f"output\t{OUTPUT / 'study_celltype_reference_correlation_heatmap.png'}")


PROGRAM_SCORE_SOURCE_COLUMNS = [
    "program",
    "program_order",
    "column_key",
    "column_order",
    "column_type",
    "column_label",
    "study",
    "study_display",
    "author_label",
    "author_label_display",
    "reference_subclass",
    "common_label",
    "mapping_status",
    "source_unit_kind",
    "aggregation_unit_kind",
    "aggregation_unit_count",
    "contributing_group_count",
    "n_cells_total",
    "region_coverage",
    "condition_coverage",
    "raw_score",
    "column_zscore",
    "raw_score_reason",
    "zscore_reason",
]


def _program_score_column_order() -> list[dict[str, object]]:
    order_path = OUTPUT / "study_celltype_reference_correlation_heatmap.tsv"
    columns = [
        "column_key",
        "study",
        "study_display",
        "study_order",
        "author_label",
        "author_label_display",
        "author_type_order",
        "column_order",
    ]
    order = pd.read_csv(order_path, sep="\t", usecols=columns, dtype=str, keep_default_na=False)
    order = order.drop_duplicates("column_key", keep="first").copy()
    order["column_order"] = pd.to_numeric(order["column_order"], errors="coerce")
    order["study_order"] = pd.to_numeric(order["study_order"], errors="coerce")
    order["author_type_order"] = pd.to_numeric(order["author_type_order"], errors="coerce")
    order = order.sort_values("column_order", kind="stable")
    return order.to_dict("records")


def _program_score_column_profile(
    records: list[dict[str, object]],
) -> tuple[np.ndarray, dict[str, object]]:
    units: dict[str, dict[str, object]] = {}
    for record in records:
        unit_key, unit_kind = _heatmap_aggregation_unit(record)
        units.setdefault(unit_key, {"unit_kind": unit_kind, "records": []})["records"].append(record)
    unit_profiles = []
    for unit in units.values():
        matrix = np.asarray(
            [
                [record.get(name, np.nan) for name in UMAP_PROGRAMS]
                for record in unit["records"]
            ],
            dtype=np.float64,
        )
        unit_profiles.append(np.median(matrix, axis=0))
    if unit_profiles:
        profile = np.median(np.asarray(unit_profiles, dtype=np.float64), axis=0)
    else:
        profile = np.full(len(UMAP_PROGRAMS), np.nan, dtype=np.float64)
    metadata = {
        "source_unit_kind": _heatmap_join_values(
            record.get("unit_kind", "") for record in records
        ),
        "aggregation_unit_kind": _heatmap_join_values(
            unit.get("unit_kind", "") for unit in units.values()
        ),
        "aggregation_unit_count": len(units),
        "contributing_group_count": len(records),
        "n_cells_total": int(sum(int(record.get("n_cells", 0)) for record in records)),
        "region_coverage": _heatmap_join_values(record.get("region", "") for record in records),
        "condition_coverage": _heatmap_join_values(
            record.get("condition", "") for record in records
        ),
        "common_label": _heatmap_join_values(
            record.get("common_label", "") for record in records
        ),
        "mapping_status": _heatmap_join_values(
            record.get("mapping_status", "") for record in records
        ),
    }
    return profile, metadata


def _program_score_column_zscore(values: np.ndarray) -> tuple[np.ndarray, str, str]:
    array = np.asarray(values, dtype=np.float64)
    if not np.isfinite(array).all():
        return np.full(array.shape, np.nan, dtype=np.float64), "NONFINITE_RAW_SCORE", "NONFINITE_RAW_SCORE"
    standard_deviation = float(np.std(array, ddof=1))
    if not np.isfinite(standard_deviation) or standard_deviation == 0.0:
        return np.full(array.shape, np.nan, dtype=np.float64), "", "ZERO_VARIANCE_COLUMN"
    return (array - float(np.mean(array))) / standard_deviation, "", ""


def run_study_celltype_program_score_source() -> None:
    profile_path = OUTPUT / "library_celltype_median_profiles.tsv"
    reference_path = OUTPUT / "reference_subclass_median_profiles.tsv"
    source_path = OUTPUT / "study_celltype_program_score_heatmap.tsv"
    external = _clean_median_profile_frame(pd.read_csv(profile_path, sep="\t"))
    reference = _clean_median_profile_frame(pd.read_csv(reference_path, sep="\t"))
    external_order = _program_score_column_order()
    external_records = external.to_dict("records")
    groups: dict[tuple[str, str], list[dict[str, object]]] = {}
    for record in external_records:
        key = (clean(record.get("study", "")), clean(record.get("author_label", "")))
        groups.setdefault(key, []).append(record)

    column_profiles: list[tuple[dict[str, object], np.ndarray, dict[str, object]]] = []
    for definition in external_order:
        key = (clean(definition.get("study", "")), clean(definition.get("author_label", "")))
        profile, metadata = _program_score_column_profile(groups.get(key, []))
        column_profiles.append((definition, profile, metadata))

    reference_profiles: list[tuple[dict[str, object], np.ndarray, dict[str, object]]] = []
    for index, record in enumerate(reference.to_dict("records"), start=1):
        label = clean(record.get("common_label", ""))
        profile = np.asarray([record.get(name, np.nan) for name in UMAP_PROGRAMS], dtype=np.float64)
        definition = {
            "column_key": f"reference_1M::{label}",
            "column_order": index,
            "column_type": "reference_subclass",
            "column_label": label,
            "study": "reference_1M",
            "study_display": "1M reference",
            "author_label": "",
            "author_label_display": "",
            "reference_subclass": label,
        }
        reference_n_cells = pd.to_numeric(record.get("n_cells", 0), errors="coerce")
        metadata = {
            "source_unit_kind": clean(record.get("unit_kind", "")),
            "aggregation_unit_kind": "reference_direct",
            "aggregation_unit_count": 1,
            "contributing_group_count": 1,
            "n_cells_total": int(reference_n_cells) if pd.notna(reference_n_cells) else 0,
            "region_coverage": "",
            "condition_coverage": "",
            "common_label": label,
            "mapping_status": clean(record.get("mapping_status", "")),
        }
        reference_profiles.append((definition, profile, metadata))

    rows = []
    zero_variance_columns = []
    nonfinite_columns = []
    all_columns = reference_profiles + column_profiles
    for global_order, (definition, profile, metadata) in enumerate(all_columns, start=1):
        zscore, raw_reason, zscore_reason = _program_score_column_zscore(profile)
        if zscore_reason == "ZERO_VARIANCE_COLUMN":
            zero_variance_columns.append(definition["column_key"])
        if raw_reason:
            nonfinite_columns.append(definition["column_key"])
        for program_order, program in enumerate(UMAP_PROGRAMS, start=1):
            raw_value = float(profile[program_order - 1]) if np.isfinite(profile[program_order - 1]) else np.nan
            z_value = float(zscore[program_order - 1]) if np.isfinite(zscore[program_order - 1]) else np.nan
            rows.append(
                {
                    "program": program,
                    "program_order": program_order,
                    "column_key": definition["column_key"],
                    "column_order": global_order,
                    "column_type": definition.get("column_type", "external_author_celltype"),
                    "column_label": definition.get("column_label", definition.get("author_label_display", "")),
                    "study": definition.get("study", ""),
                    "study_display": definition.get("study_display", ""),
                    "author_label": definition.get("author_label", ""),
                    "author_label_display": definition.get("author_label_display", ""),
                    "reference_subclass": definition.get("reference_subclass", ""),
                    "common_label": metadata.get("common_label", ""),
                    "mapping_status": metadata.get("mapping_status", ""),
                    "source_unit_kind": metadata.get("source_unit_kind", ""),
                    "aggregation_unit_kind": metadata.get("aggregation_unit_kind", ""),
                    "aggregation_unit_count": metadata.get("aggregation_unit_count", 0),
                    "contributing_group_count": metadata.get("contributing_group_count", 0),
                    "n_cells_total": metadata.get("n_cells_total", 0),
                    "region_coverage": metadata.get("region_coverage", ""),
                    "condition_coverage": metadata.get("condition_coverage", ""),
                    "raw_score": raw_value,
                    "column_zscore": z_value,
                    "raw_score_reason": raw_reason,
                    "zscore_reason": zscore_reason,
                }
            )
    table = pd.DataFrame(rows, columns=PROGRAM_SCORE_SOURCE_COLUMNS)
    table.to_csv(source_path, sep="\t", index=False, na_rep="NA", float_format="%.17g")
    print("PROGRAM_SCORE_SOURCE_COMPLETE")
    print(f"programs\t{len(UMAP_PROGRAMS)}")
    print(f"reference_columns\t{len(reference_profiles)}")
    print(f"external_columns\t{len(column_profiles)}")
    print(f"columns\t{len(all_columns)}")
    print(f"matrix_cells\t{len(UMAP_PROGRAMS) * len(all_columns)}")
    print(f"source_rows\t{len(table)}")
    print(f"zero_variance_columns\t{len(zero_variance_columns)}")
    for column_key in zero_variance_columns:
        print(f"zero_variance_column\t{column_key}")
    print(f"nonfinite_columns\t{len(nonfinite_columns)}")
    for column_key in nonfinite_columns:
        print(f"nonfinite_column\t{column_key}")
    print(f"output\t{source_path}")


PROGRAM_SCORE_ANNOTATION_COLUMNS = [
    "program_dominant_class_raw",
    "program_dominant_class",
    "program_dominant_subclass",
    "program_confidence",
    "program_annotation_source",
    "column_major_class",
    "column_annotation_raw",
    "column_annotation_status",
    "column_annotation_source",
]


def _f1b_program_class(value: object) -> str:
    key = clean(value).casefold()
    if key in {"exc", "excitatory", "glutamatergic"}:
        return "excitatory"
    if key in {"inh", "inhibitory", "gabaergic"}:
        return "inhibitory"
    if key in {"glia", "nonneuron", "non-neuronal", "nonneuronal", "vascular"}:
        return "non-neuronal"
    return "Unknown"


def _h5ad_original_class(value: object) -> str:
    key = clean(value).casefold().replace("_", "-")
    if "glutamatergic" in key:
        return "excitatory"
    if "gabaergic" in key:
        return "inhibitory"
    if any(
        token in key
        for token in (
            "non-neuronal",
            "oligodendro",
            "astro",
            "opc",
            "microgl",
            "immune",
            "endothel",
            "vlmc",
            "pericyte",
        )
    ):
        return "non-neuronal"
    return ""


def _crosswalk_original_class(value: object) -> str:
    key = clean(value).casefold()
    if key == "excitatory neuron":
        return "excitatory"
    if key == "inhibitory neuron":
        return "inhibitory"
    if key in {
        "astrocyte",
        "oligodendrocyte",
        "opc",
        "endothelial",
        "microglia/immune",
        "glia",
        "non-neuronal",
        "vascular",
    }:
        return "non-neuronal"
    if "/" in key or any(token in key for token in ("mixed", "other", "unspecified", "author cluster", "t cell", "fibroblast", "pericyte")):
        return "Mixed"
    return ""


def _accumulate_annotation(
    target: dict[str, dict[str, set[str]]],
    label: object,
    broad_class: str,
    raw_value: object,
    source: object,
) -> None:
    key = clean(label).casefold()
    if not key:
        return
    entry = target.setdefault(key, {"classes": set(), "raw": set(), "sources": set()})
    if broad_class:
        entry["classes"].add(broad_class)
    raw = clean(raw_value)
    if raw:
        entry["raw"].add(raw)
    source_text = clean(source)
    if source_text:
        entry["sources"].add(source_text)


def _read_h5ad_class_annotations(
    paths: Iterable[Path],
    label_field_candidates: tuple[str, ...],
    class_field_candidates: tuple[str, ...],
) -> dict[str, dict[str, set[str]]]:
    target: dict[str, dict[str, set[str]]] = {}
    for path in paths:
        source = ad.read_h5ad(path, backed="r")
        try:
            label_field = next(
                (field for field in label_field_candidates if field in source.obs.columns),
                "",
            )
            class_field = next(
                (field for field in class_field_candidates if field in source.obs.columns),
                "",
            )
            if not label_field or not class_field:
                continue
            labels = source.obs[label_field].astype(str).tolist()
            classes = source.obs[class_field].astype(str).tolist()
            for label, raw_class in zip(labels, classes):
                _accumulate_annotation(
                    target,
                    label,
                    _h5ad_original_class(raw_class),
                    raw_class,
                    f"{path}::obs[{class_field}]",
                )
        finally:
            source.file.close()
    return target


def _read_crosswalk_class_annotations(study: str) -> dict[str, dict[str, set[str]]]:
    target: dict[str, dict[str, set[str]]] = {}
    path = INPUT / study / "backup/author_label_to_canonical22_crosswalk.tsv"
    table = read_mapping_table(path, "dataset\t")
    for record in table.to_dict("records"):
        author = clean(record.get("author_label", ""))
        raw_class = clean(record.get("broad_class", ""))
        _accumulate_annotation(
            target,
            author,
            _crosswalk_original_class(raw_class),
            raw_class,
            str(path),
        )
    return target


def _resolve_original_class(
    entry: dict[str, set[str]] | None,
    label: str,
) -> tuple[str, str, str, str]:
    if not label:
        return "Unknown", "", "MISSING_ORIGINAL_ANNOTATION", ""
    if not entry or not entry["classes"]:
        return "Unknown", ";".join(sorted(entry["raw"])) if entry else "", "UNKNOWN_ORIGINAL_CLASS", ";".join(sorted(entry["sources"])) if entry else ""
    classes = entry["classes"]
    if len(classes) == 1:
        broad_class = next(iter(classes))
        status = "AUTHORITATIVE_ORIGINAL_CLASS" if broad_class in {"excitatory", "inhibitory", "non-neuronal"} else "UNKNOWN_OR_MIXED_ORIGINAL_CLASS"
    else:
        broad_class = "Mixed"
        status = "CONFLICTING_ORIGINAL_CLASS"
    return broad_class, ";".join(sorted(entry["raw"])), status, ";".join(sorted(entry["sources"]))


def run_annotate_study_celltype_program_source() -> None:
    source_path = OUTPUT / "study_celltype_program_score_heatmap.tsv"
    annotation_path = Path((__import__("os").environ["NMF_WORK_ROOT"] + "/tables/TableS3_program_annotation.tsv"))
    reference_class_path = PROJECT / "archived/figures/fig1/_intermediate/subclass_class.csv"
    with source_path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        fieldnames = list(reader.fieldnames or [])
        records = list(reader)
    for field in PROGRAM_SCORE_ANNOTATION_COLUMNS:
        if field not in fieldnames:
            fieldnames.append(field)

    program_table = pd.read_csv(annotation_path, sep="\t", dtype=str).fillna("")
    program_annotations = {
        clean(record.get("new_P", "")): record
        for record in program_table.to_dict("records")
    }
    subclass_table = pd.read_csv(reference_class_path, dtype=str).fillna("")
    reference_classes = {
        clean(record.get("subclass", "")).casefold(): record
        for record in subclass_table.to_dict("records")
    }

    allen_annotations = _read_h5ad_class_annotations(
        [ALLEN_H5AD],
        ("subclass", "Subclass"),
        ("class", "Class"),
    )
    seaad_annotations = _read_h5ad_class_annotations(
        [
            SEAAD_META / "SEAAD_DLPFC_reconstructed_raw_counts.h5ad",
            SEAAD_META / "SEAAD_MTG_reconstructed_raw_counts.h5ad",
        ],
        ("Subclass", "subclass"),
        ("cell_type", "Cell.Type", "class", "Class"),
    )
    study_annotations = {
        study: _read_crosswalk_class_annotations(study)
        for study in (
            "Grubman_2019_GSE138852",
            "Lake_2018_GSE97930",
            "Schirmer_2019_MS_UCSC",
            "SingleSoma_AD_PFC_CELLxGENE",
            "Tran_2021_reward_cortex_CELLxGENE",
            "GSE144136",
            "GSE174367",
            "GSE291605",
        )
    }

    class_counts: dict[str, int] = {}
    for record in records:
        program = clean(record.get("program", ""))
        program_annotation = program_annotations.get(program, {})
        raw_program_class = clean(program_annotation.get("dominant_class", ""))
        record["program_dominant_class_raw"] = raw_program_class
        record["program_dominant_class"] = _f1b_program_class(raw_program_class)
        record["program_dominant_subclass"] = clean(program_annotation.get("dominant_subclass", ""))
        record["program_confidence"] = clean(program_annotation.get("confidence", ""))
        record["program_annotation_source"] = str(annotation_path)

        column_type = clean(record.get("column_type", ""))
        if column_type == "reference_subclass":
            label = clean(record.get("reference_subclass", ""))
            reference_record = reference_classes.get(label.casefold(), {})
            raw_class = clean(reference_record.get("class", ""))
            broad_class = _f1b_program_class(raw_class)
            status = "AUTHORITATIVE_REFERENCE_CLASS" if broad_class != "Unknown" else "UNKNOWN_REFERENCE_CLASS"
            source = str(reference_class_path)
        else:
            study = clean(record.get("study", ""))
            label = clean(record.get("author_label", ""))
            if study == "Allen":
                annotation_entry = allen_annotations.get(label.casefold())
            elif study == "SEAAD":
                annotation_entry = seaad_annotations.get(label.casefold())
            else:
                annotation_entry = study_annotations.get(study, {}).get(label.casefold())
            broad_class, raw_class, status, source = _resolve_original_class(annotation_entry, label)
        record["column_major_class"] = broad_class
        record["column_annotation_raw"] = raw_class
        record["column_annotation_status"] = status
        record["column_annotation_source"] = source
        class_counts[broad_class] = class_counts.get(broad_class, 0) + 1

    with source_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(records)
    print("PROGRAM_SCORE_SOURCE_ANNOTATED")
    print(f"source_rows\t{len(records)}")
    print(f"columns\t{len(set(record.get('column_key', '') for record in records))}")
    for broad_class in sorted(class_counts):
        print(f"column_annotation_class\t{broad_class}\t{class_counts[broad_class]}")
    print(f"output\t{source_path}")


def _read_h5ad_label_class_map(
    paths: Iterable[Path],
    label_field_candidates: tuple[str, ...],
    class_field_candidates: tuple[str, ...],
    lineage_field: bool = False,
) -> dict[str, str]:
    assignments: dict[str, set[str]] = {}
    for path in paths:
        source = ad.read_h5ad(path, backed="r")
        try:
            label_field = next(
                (field for field in label_field_candidates if field in source.obs.columns),
                "",
            )
            class_field = next(
                (field for field in class_field_candidates if field in source.obs.columns),
                "",
            )
            if not label_field or not class_field:
                continue
            table = source.obs[[label_field, class_field]].drop_duplicates()
            for label, raw_class in table.itertuples(index=False, name=None):
                label = clean(label)
                raw_class = clean(raw_class)
                if not label:
                    continue
                if lineage_field:
                    broad_class = _native_lineage_class(raw_class)
                else:
                    broad_class = _native_class_class(raw_class)
                if broad_class:
                    assignments.setdefault(label.casefold(), set()).add(broad_class)
        finally:
            source.file.close()
    return {
        label: next(iter(classes)) if len(classes) == 1 else "Mixed"
        for label, classes in assignments.items()
    }


def _native_class_class(value: object) -> str:
    key = clean(value).casefold()
    if "glutamatergic" in key:
        return "Exc"
    if "gabaergic" in key:
        return "Inh"
    if "non-neuronal" in key or "nonneuronal" in key:
        return ""
    return ""


def _native_lineage_class(value: object) -> str:
    key = clean(value).casefold().replace("_", "-")
    if "glutamatergic" in key:
        return "Exc"
    if "gabaergic" in key:
        return "Inh"
    if "astro" in key:
        return "Ast"
    if "oligodendrocyte precursor" in key or "opc" in key:
        return "OPC"
    if "oligodendro" in key:
        return "Oligo"
    if "microgl" in key:
        return "Micro"
    if "endothel" in key:
        return "Endo"
    if "vascular leptomeningeal" in key or "vlmc" in key:
        return "VLMC"
    return ""


def _subclass_to_broad_class(label: object, class_value: object) -> str:
    key = clean(label).casefold()
    direct = {
        "ast": "Ast",
        "astrocyte": "Ast",
        "oligo": "Oligo",
        "oligodendrocyte": "Oligo",
        "opc": "OPC",
        "opcs": "OPC",
        "micro": "Micro",
        "microglia": "Micro",
        "endo": "Endo",
        "endothelial": "Endo",
        "vlmc": "VLMC",
    }
    if key in direct:
        return direct[key]
    class_key = clean(class_value).casefold()
    if class_key == "excitatory":
        return "Exc"
    if class_key == "inhibitory":
        return "Inh"
    return ""


def _crosswalk_broad_class_for_scatter(raw_class: object, author_label: object) -> str:
    key = clean(raw_class).casefold()
    author = clean(author_label).casefold()
    if key == "excitatory neuron":
        return "Exc"
    if key == "inhibitory neuron":
        return "Inh"
    if key == "astrocyte":
        return "Ast"
    if key == "oligodendrocyte":
        return "Oligo"
    if key == "opc":
        return "OPC"
    if key == "endothelial":
        return "Endo"
    if key == "microglia/immune" and author in {"micro", "microglia", "mg"}:
        return "Micro"
    if key == "vlmc" and author in {"vlmc", "mural"}:
        return "VLMC"
    return ""


def _read_scatter_library_lookup(
    path: Path, field: str, id_field: str | None = None
) -> dict[str, str]:
    source = ad.read_h5ad(path, backed="r")
    try:
        values = [clean(value) for value in source.obs[field].tolist()]
        ids = [clean(value) for value in source.obs_names]
        lookup = dict(zip(ids, values))
        if id_field and id_field in source.obs.columns:
            for cell_id, value in zip(source.obs[id_field].tolist(), values):
                cell_id = clean(cell_id)
                if cell_id:
                    lookup[cell_id] = value
        return lookup
    finally:
        source.file.close()


def _scatter_external_broad_class_map(study: str) -> dict[str, str]:
    if study == "Allen":
        path = ALLEN_H5AD
        source = ad.read_h5ad(path, backed="r")
        try:
            fields = [field for field in ("class", "subclass") if field in source.obs.columns]
            if len(fields) != 2:
                return {}
            table = source.obs[fields].drop_duplicates()
            assignments: dict[str, set[str]] = {}
            for raw_class, label in table.itertuples(index=False, name=None):
                broad = _native_class_class(raw_class)
                if not broad:
                    broad = _subclass_to_broad_class(label, raw_class)
                if broad:
                    assignments.setdefault(clean(label).casefold(), set()).add(broad)
            return {
                label: next(iter(classes)) if len(classes) == 1 else "Mixed"
                for label, classes in assignments.items()
            }
        finally:
            source.file.close()
    if study == "SEAAD":
        lineage_map = _read_h5ad_label_class_map(
            [
                SEAAD_META / "SEAAD_DLPFC_reconstructed_raw_counts.h5ad",
                SEAAD_META / "SEAAD_MTG_reconstructed_raw_counts.h5ad",
            ],
            ("Subclass", "subclass"),
            ("cell_type", "Cell.Type", "class", "Class"),
            lineage_field=True,
        )
        return lineage_map
    path = INPUT / study / "backup/author_label_to_canonical22_crosswalk.tsv"
    table = read_mapping_table(path, "dataset\t")
    assignments: dict[str, set[str]] = {}
    for record in table.to_dict("records"):
        label = clean(record.get("author_label", ""))
        broad = _crosswalk_broad_class_for_scatter(
            record.get("broad_class", ""), label
        )
        if label and broad:
            assignments.setdefault(label.casefold(), set()).add(broad)
    return {
        label: next(iter(classes)) if len(classes) == 1 else "Mixed"
        for label, classes in assignments.items()
    }


def _scatter_external_frame(study: str) -> pd.DataFrame:
    path = OUTPUT / f"{study}.cell_scores.parquet"
    columns = [
        "cell_id",
        "source_unit",
        "donor",
        "region",
        "condition",
        "author_label",
        *UMAP_PROGRAMS,
    ]
    frame = pd.read_parquet(path, columns=columns)
    for column in ["cell_id", "source_unit", "donor", "region", "condition", "author_label"]:
        frame[column] = frame[column].map(clean)
    broad_map = _scatter_external_broad_class_map(study)
    frame["broad_class"] = frame["author_label"].map(
        lambda value: broad_map.get(clean(value).casefold(), "")
    )
    if study == "Allen":
        library_lookup = _read_scatter_library_lookup(ALLEN_H5AD, "Specimen ID")
        frame["library"] = frame["cell_id"].map(library_lookup).fillna("").map(clean)
        frame["unit_kind"] = "sample_specimen"
        frame["library_is_proxy"] = False
    elif study == "SEAAD":
        frame["library"] = frame["source_unit"].map(clean)
        frame["unit_kind"] = "source_unit_library"
        frame["library_is_proxy"] = False
    elif study == "Schirmer_2019_MS_UCSC":
        frame["library"] = frame["source_unit"].map(clean)
        frame["unit_kind"] = "source_unit_library"
        frame["library_is_proxy"] = False
    elif study == "Tran_2021_reward_cortex_CELLxGENE":
        frame["library"] = frame["source_unit"].map(clean)
        frame["unit_kind"] = "donor_region_sample_proxy"
        frame["library_is_proxy"] = True
    else:
        specs = {
            "SingleSoma_AD_PFC_CELLxGENE": (SINGLESOMA_H5AD, "Sample.ID", "sample_id"),
            "GSE144136": (GENERIC_H5ADS["GSE144136"], "sample_id", "sample_id"),
            "GSE174367": (GENERIC_H5ADS["GSE174367"], "sample_id", "sample_id"),
            "GSE291605": (GENERIC_H5ADS["GSE291605"], "library_name", "library_name"),
        }
        path, field, unit_kind = specs[study]
        lookup = _read_scatter_library_lookup(path, field)
        frame["library"] = frame["cell_id"].map(lookup).fillna("").map(clean)
        frame["unit_kind"] = unit_kind
        frame["library_is_proxy"] = False
    return frame


def _scatter_aggregate_profiles(
    frame: pd.DataFrame,
    value_columns: list[str],
    source_unit_kind: str | None = None,
    allowed_classes: set[str] | None = None,
) -> tuple[dict[str, np.ndarray], dict[str, dict[str, object]]]:
    frame = frame.copy()
    accepted_classes = set(BROAD_CONSISTENCY_CLASSES) if allowed_classes is None else set(allowed_classes)
    frame = frame[
        frame["broad_class"].isin(accepted_classes)
        & frame["library"].astype(str).ne("")
        & frame["donor"].astype(str).ne("")
    ].copy()
    library_keys = ["library", "donor", "broad_class", "region", "condition"]
    library_profiles = []
    for key, positions in frame.groupby(library_keys, sort=False, dropna=False).indices.items():
        subset = frame.iloc[positions]
        profile = np.median(subset[value_columns].to_numpy(dtype=np.float64), axis=0)
        library, donor, broad_class, region, condition = key
        library_profiles.append(
            {
                "library": library,
                "donor": donor,
                "broad_class": broad_class,
                "region": region,
                "condition": condition,
                "profile": profile,
                "n_cells": int(len(subset)),
                "unit_kind": source_unit_kind or _heatmap_join_values(subset.get("unit_kind", "")),
                "library_is_proxy": bool(subset.get("library_is_proxy", pd.Series([False])).any()),
            }
        )
    if not library_profiles:
        return {}, {}
    library_frame = pd.DataFrame(library_profiles)
    profiles: dict[str, np.ndarray] = {}
    metadata: dict[str, dict[str, object]] = {}
    for broad_class, positions in library_frame.groupby("broad_class", sort=False).indices.items():
        subset = library_frame.iloc[positions]
        donor_profiles = []
        for donor, donor_positions in subset.groupby("donor", sort=False).indices.items():
            donor_subset = subset.iloc[donor_positions]
            donor_profiles.append(
                {
                    "donor": donor,
                    "profile": np.median(
                        np.stack(donor_subset["profile"].to_numpy()), axis=0
                    ),
                }
            )
        final_profile = np.median(
            np.stack([entry["profile"] for entry in donor_profiles]), axis=0
        )
        profiles[broad_class] = final_profile
        metadata[broad_class] = {
            "n_cells": int(frame.loc[frame["broad_class"].eq(broad_class), "n_cells"].sum())
            if "n_cells" in frame.columns
            else int(sum(entry["n_cells"] for entry in library_profiles if entry["broad_class"] == broad_class)),
            "library_count": int(subset["library"].nunique()),
            "donor_count": int(len(donor_profiles)),
            "library_group_count": int(len(subset)),
            "region_coverage": _heatmap_join_values(subset["region"]),
            "condition_coverage": _heatmap_join_values(subset["condition"]),
            "unit_kind": _heatmap_join_values(subset["unit_kind"]),
            "library_is_proxy": bool(subset["library_is_proxy"].any()),
        }
    return profiles, metadata


def _scatter_reference_profiles(
    neuron_only: bool = False,
) -> tuple[dict[str, np.ndarray], dict[str, dict[str, object]]]:
    old_columns = _load_umap_retained_columns()
    source_columns = [str(value) for value in range(1, 61)]
    scores = pd.read_parquet(DISCOVERY_PROFILE_PATH, columns=source_columns)
    cell_ids = [clean(value) for value in scores.index.tolist()]
    obs = pd.read_csv(
        DISCOVERY_OBS_PATH,
        index_col=0,
        usecols=["Unnamed: 0", "donor", "region", "subclass", "library_prep"],
        dtype=str,
    ).fillna("")
    obs.index = [clean(value) for value in obs.index]
    obs = obs.reindex(cell_ids)
    subclass_class_table = pd.read_csv(
        PROJECT / "archived/figures/fig1/_intermediate/subclass_class.csv",
        dtype=str,
    ).fillna("")
    subclass_class = {
        clean(record.get("subclass", "")): clean(record.get("class", ""))
        for record in subclass_class_table.to_dict("records")
    }
    full_values = scores[source_columns].to_numpy(dtype=np.float64)
    row_sum = full_values.sum(axis=1, keepdims=True)
    normalized = np.divide(
        full_values,
        row_sum,
        out=np.zeros_like(full_values),
        where=row_sum > 0.0,
    )
    selected = normalized[:, np.asarray(old_columns, dtype=np.int64) - 1]
    frame = pd.DataFrame(selected, columns=UMAP_PROGRAMS)
    frame["library"] = obs["library_prep"].astype(str).map(clean).to_numpy()
    frame["donor"] = obs["donor"].astype(str).map(clean).to_numpy()
    frame["region"] = obs["region"].astype(str).map(clean).to_numpy()
    frame["condition"] = ""
    frame["subclass"] = obs["subclass"].astype(str).map(clean).to_numpy()
    frame["broad_class"] = [
        _subclass_to_broad_class(label, subclass_class.get(label, ""))
        for label in frame["subclass"]
    ]
    frame["unit_kind"] = "reference_library_prep"
    frame["library_is_proxy"] = False
    if neuron_only:
        frame = frame[frame["broad_class"].isin({"Exc", "Inh"})].copy()
        frame["broad_class"] = "Neuron"
        return _scatter_aggregate_profiles(
            frame,
            UMAP_PROGRAMS,
            source_unit_kind="reference_library_prep",
            allowed_classes={"Neuron"},
        )
    return _scatter_aggregate_profiles(frame, UMAP_PROGRAMS, source_unit_kind="reference_library_prep")


def _scatter_cosine_similarity(
    reference: np.ndarray,
    external: np.ndarray,
) -> tuple[float, int, str]:
    reference = np.asarray(reference, dtype=np.float64)
    external = np.asarray(external, dtype=np.float64)
    paired = np.isfinite(reference) & np.isfinite(external)
    n_pairs = int(paired.sum())
    if n_pairs == 0:
        return float("nan"), n_pairs, "no_finite_pairs"
    if n_pairs != len(reference) or n_pairs != len(external):
        return float("nan"), n_pairs, "nonfinite_pairs"
    reference = reference[paired]
    external = external[paired]
    reference_scale = float(np.max(np.abs(reference)))
    external_scale = float(np.max(np.abs(external)))
    if reference_scale == 0.0 and external_scale == 0.0:
        return float("nan"), n_pairs, "zero_reference_and_external_norm"
    if reference_scale == 0.0:
        return float("nan"), n_pairs, "zero_reference_norm"
    if external_scale == 0.0:
        return float("nan"), n_pairs, "zero_external_norm"
    reference = reference / reference_scale
    external = external / external_scale
    denominator = float(np.linalg.norm(reference) * np.linalg.norm(external))
    if not np.isfinite(denominator) or denominator == 0.0:
        return float("nan"), n_pairs, "zero_norm"
    value = float(np.dot(reference, external) / denominator)
    if not np.isfinite(value):
        return float("nan"), n_pairs, "undefined"
    return value, n_pairs, ""


def run_neuron_program_consistency() -> None:
    reference_profiles, _ = _scatter_reference_profiles(neuron_only=True)
    reference = reference_profiles.get("Neuron")
    if reference is None:
        raise RuntimeError("reference Neuron profile is unavailable")
    rows = []
    for study in BROAD_CONSISTENCY_STUDIES:
        frame = _scatter_external_frame(study)
        native_members = [
            value
            for value in ("Exc", "Inh")
            if value in set(frame["broad_class"].astype(str))
        ]
        neuron_frame = frame[frame["broad_class"].isin({"Exc", "Inh"})].copy()
        neuron_frame["broad_class"] = "Neuron"
        profiles, _ = _scatter_aggregate_profiles(
            neuron_frame,
            UMAP_PROGRAMS,
            allowed_classes={"Neuron"},
        )
        external = profiles.get("Neuron")
        if external is None:
            cosine, n_pairs, na_reason = float("nan"), 0, "no_neuron_profile"
        else:
            cosine, n_pairs, na_reason = _scatter_cosine_similarity(reference, external)
        if native_members == ["Exc", "Inh"]:
            coverage_note = "Exc+Inh"
        elif native_members == ["Inh"]:
            coverage_note = "Inh-only; Exc absent"
        elif native_members == ["Exc"]:
            coverage_note = "Exc-only; Inh absent"
        else:
            coverage_note = "no_Exc_or_Inh"
        rows.append(
            {
                "study": study,
                "cosine_similarity": cosine,
                "n_pairs": n_pairs,
                "neuron_members": "+".join(native_members),
                "coverage_note": coverage_note,
                "na_reason": na_reason,
            }
        )
    table = pd.DataFrame(
        rows,
        columns=[
            "study",
            "cosine_similarity",
            "n_pairs",
            "neuron_members",
            "coverage_note",
            "na_reason",
        ],
    )
    table.to_csv(NEURON_OUTPUT, sep="\t", index=False)
    print("NEURON_PROGRAM_CONSISTENCY_COMPLETE")
    print(f"rows\t{len(table)}")
    print(f"output\t{NEURON_OUTPUT}")
    for row in rows:
        cosine_text = "NA" if not np.isfinite(row["cosine_similarity"]) else f"{row['cosine_similarity']:.8f}"
        reason_text = row["na_reason"] or "OK"
        print(
            f"neuron_cosine\t{row['study']}\t{cosine_text}\t{row['n_pairs']}\t"
            f"{row['neuron_members']}\t{row['coverage_note']}\t{reason_text}"
        )


def _cell_nearest_reference_vectors() -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    old_columns = np.asarray(_load_umap_retained_columns(), dtype=np.int64) - 1
    source_columns = [str(value) for value in range(1, 61)]
    scores = pd.read_parquet(DISCOVERY_PROFILE_PATH, columns=source_columns)
    cell_ids = np.asarray([clean(value) for value in scores.index.tolist()], dtype=str)
    obs = pd.read_csv(
        DISCOVERY_OBS_PATH,
        index_col=0,
        usecols=["Unnamed: 0", "subclass"],
        dtype=str,
    ).fillna("")
    obs.index = [clean(value) for value in obs.index]
    subclasses = obs.reindex(cell_ids)["subclass"].astype(str).map(clean).to_numpy(dtype=str)
    subclass_class_table = pd.read_csv(
        PROJECT / "archived/figures/fig1/_intermediate/subclass_class.csv",
        dtype=str,
    ).fillna("")
    subclass_class = {
        clean(record.get("subclass", "")): clean(record.get("class", ""))
        for record in subclass_class_table.to_dict("records")
    }
    broad_classes = np.asarray(
        [
            _subclass_to_broad_class(label, subclass_class.get(label, ""))
            for label in subclasses
        ],
        dtype=str,
    )
    full_values = scores[source_columns].to_numpy(dtype=np.float64)
    row_sum = full_values.sum(axis=1, keepdims=True)
    normalized = np.divide(
        full_values,
        row_sum,
        out=np.zeros_like(full_values),
        where=row_sum > 0.0,
    )
    selected = normalized[:, old_columns]
    finite = np.isfinite(selected).all(axis=1)
    norms = np.zeros(selected.shape[0], dtype=np.float64)
    if np.any(finite):
        norms[finite] = np.linalg.norm(selected[finite], axis=1)
    valid = finite & (norms > 0.0)
    coordinates = selected[valid] / norms[valid, None]
    return coordinates, cell_ids[valid], subclasses[valid], broad_classes[valid]


def run_cell_nearest_reference() -> None:
    import torch

    if not torch.cuda.is_available():
        raise RuntimeError("CUDA is unavailable; exact nearest-reference production requires the approved GPU route")
    torch.backends.cuda.matmul.allow_tf32 = False
    device = torch.device("cuda")
    reference_coordinates, reference_cell_ids, reference_subclasses, reference_broad_classes = _cell_nearest_reference_vectors()
    reference_tensor = torch.as_tensor(
        reference_coordinates,
        dtype=torch.float64,
        device=device,
    )
    output_columns = [
        "cell_id",
        "study",
        "author_label",
        "broad_class",
        "donor",
        "library",
        "source_unit",
        "region",
        "condition",
        "unit_kind",
        "library_is_proxy",
        "nearest_reference_cell_id",
        "nearest_reference_subclass",
        "nearest_reference_broad_class",
        "cosine_similarity",
        "na_reason",
    ]
    for study in BROAD_CONSISTENCY_STUDIES:
        frame = _scatter_external_frame(study)
        frame = frame[frame["broad_class"].isin(BROAD_CONSISTENCY_CLASSES)].copy()
        output_path = OUTPUT / f"{CELL_NEAREST_OUTPUT_PREFIX}{study}.parquet"
        writer = None
        rows_written = 0
        valid_rows = 0
        na_rows = 0
        try:
            for start in range(0, len(frame), CELL_NEAREST_QUERY_BATCH_SIZE):
                stop = min(len(frame), start + CELL_NEAREST_QUERY_BATCH_SIZE)
                chunk = frame.iloc[start:stop].copy()
                values = chunk[UMAP_PROGRAMS].to_numpy(dtype=np.float64)
                finite = np.isfinite(values).all(axis=1)
                norms = np.zeros(values.shape[0], dtype=np.float64)
                if np.any(finite):
                    norms[finite] = np.linalg.norm(values[finite], axis=1)
                valid = finite & (norms > 0.0)
                nearest_indices = np.full(values.shape[0], -1, dtype=np.int64)
                cosine_values = np.full(values.shape[0], np.nan, dtype=np.float64)
                reasons = np.full(values.shape[0], "", dtype=object)
                reasons[~finite] = "nonfinite_query"
                reasons[finite & (norms == 0.0)] = "zero_query_norm"
                if np.any(valid):
                    query_coordinates = values[valid] / norms[valid, None]
                    query_tensor = torch.as_tensor(
                        query_coordinates,
                        dtype=torch.float64,
                        device=device,
                    )
                    best_values = torch.full(
                        (query_tensor.shape[0],),
                        -torch.inf,
                        dtype=torch.float64,
                        device=device,
                    )
                    best_indices = torch.full(
                        (query_tensor.shape[0],),
                        -1,
                        dtype=torch.int64,
                        device=device,
                    )
                    with torch.no_grad():
                        for ref_start in range(0, reference_tensor.shape[0], CELL_NEAREST_REFERENCE_BLOCK_SIZE):
                            ref_stop = min(reference_tensor.shape[0], ref_start + CELL_NEAREST_REFERENCE_BLOCK_SIZE)
                            similarity = query_tensor @ reference_tensor[ref_start:ref_stop].transpose(0, 1)
                            block_values, block_indices = torch.max(similarity, dim=1)
                            block_indices = block_indices + ref_start
                            better = block_values > best_values
                            best_values = torch.where(better, block_values, best_values)
                            best_indices = torch.where(better, block_indices, best_indices)
                            del similarity, block_values, block_indices, better
                    valid_positions = np.flatnonzero(valid)
                    best_indices_cpu = best_indices.detach().cpu().numpy()
                    best_values_cpu = best_values.detach().cpu().numpy()
                    nearest_indices[valid_positions] = best_indices_cpu
                    cosine_values[valid_positions] = best_values_cpu
                    no_reference = valid_positions[best_indices_cpu < 0]
                    if no_reference.size:
                        reasons[no_reference] = "no_reference_candidates"
                    del query_tensor, best_values, best_indices
                nearest_reference_cell = np.full(values.shape[0], "", dtype=object)
                nearest_reference_subclass = np.full(values.shape[0], "", dtype=object)
                nearest_reference_broad = np.full(values.shape[0], "", dtype=object)
                has_reference = nearest_indices >= 0
                nearest_reference_cell[has_reference] = reference_cell_ids[nearest_indices[has_reference]]
                nearest_reference_subclass[has_reference] = reference_subclasses[nearest_indices[has_reference]]
                nearest_reference_broad[has_reference] = reference_broad_classes[nearest_indices[has_reference]]
                out = pd.DataFrame(
                    {
                        "cell_id": chunk["cell_id"].to_numpy(dtype=str),
                        "study": study,
                        "author_label": chunk["author_label"].to_numpy(dtype=str),
                        "broad_class": chunk["broad_class"].to_numpy(dtype=str),
                        "donor": chunk["donor"].to_numpy(dtype=str),
                        "library": chunk["library"].to_numpy(dtype=str),
                        "source_unit": chunk["source_unit"].to_numpy(dtype=str),
                        "region": chunk["region"].to_numpy(dtype=str),
                        "condition": chunk["condition"].to_numpy(dtype=str),
                        "unit_kind": chunk["unit_kind"].to_numpy(dtype=str),
                        "library_is_proxy": chunk["library_is_proxy"].to_numpy(dtype=bool),
                        "nearest_reference_cell_id": nearest_reference_cell,
                        "nearest_reference_subclass": nearest_reference_subclass,
                        "nearest_reference_broad_class": nearest_reference_broad,
                        "cosine_similarity": cosine_values,
                        "na_reason": reasons,
                    },
                    columns=output_columns,
                )
                table = pa.Table.from_pandas(out, preserve_index=False)
                if writer is None:
                    writer = pq.ParquetWriter(output_path, table.schema, compression="zstd")
                writer.write_table(table)
                rows_written += len(out)
                valid_rows += int(np.isfinite(cosine_values).sum())
                na_rows += int((~np.isfinite(cosine_values)).sum())
                del chunk, values, finite, norms, valid, out, table
            if writer is None:
                empty = pd.DataFrame(
                    {
                        "cell_id": pd.Series(dtype=str),
                        "study": pd.Series(dtype=str),
                        "author_label": pd.Series(dtype=str),
                        "broad_class": pd.Series(dtype=str),
                        "donor": pd.Series(dtype=str),
                        "library": pd.Series(dtype=str),
                        "source_unit": pd.Series(dtype=str),
                        "region": pd.Series(dtype=str),
                        "condition": pd.Series(dtype=str),
                        "unit_kind": pd.Series(dtype=str),
                        "library_is_proxy": pd.Series(dtype=bool),
                        "nearest_reference_cell_id": pd.Series(dtype=str),
                        "nearest_reference_subclass": pd.Series(dtype=str),
                        "nearest_reference_broad_class": pd.Series(dtype=str),
                        "cosine_similarity": pd.Series(dtype=np.float64),
                        "na_reason": pd.Series(dtype=str),
                    },
                    columns=output_columns,
                )
                writer = pq.ParquetWriter(
                    output_path,
                    pa.Table.from_pandas(empty, preserve_index=False).schema,
                    compression="zstd",
                )
            print(
                f"study_nearest_complete\\t{study}\\trows\\t{rows_written}\\tvalid\\t{valid_rows}\\tNA\\t{na_rows}\\toutput\\t{output_path}"
            )
        finally:
            if writer is not None:
                writer.close()
        del frame
    del reference_tensor, reference_coordinates
    print("CELL_NEAREST_REFERENCE_COMPLETE")


def run_broad_program_consistency_source() -> None:
    output_path = OUTPUT / "broad_program_consistency_scatter.tsv"
    reference_profiles, reference_metadata = _scatter_reference_profiles()
    rows = []
    coverage = {}
    for study in BROAD_CONSISTENCY_STUDIES:
        frame = _scatter_external_frame(study)
        profiles, metadata = _scatter_aggregate_profiles(frame, UMAP_PROGRAMS)
        coverage[study] = {
            "groups": int(len(frame)),
            "assigned_cells": int(frame[frame["broad_class"].isin(BROAD_CONSISTENCY_CLASSES)].shape[0]),
            "classes": sorted(profiles),
            "profiles": profiles,
            "metadata": metadata,
        }
        for broad_class in sorted(set(reference_profiles) & set(profiles)):
            x_meta = reference_metadata[broad_class]
            y_meta = metadata[broad_class]
            for program in UMAP_PROGRAMS:
                x_value = float(reference_profiles[broad_class][int(program[1:]) - 1])
                y_value = float(profiles[broad_class][int(program[1:]) - 1])
                status = "OK" if np.isfinite(x_value) and np.isfinite(y_value) else "NONFINITE_SCORE"
                rows.append(
                    {
                        "point_id": f"{study}::{broad_class}::{program}",
                        "study": study,
                        "broad_class": broad_class,
                        "program": program,
                        "x_reference_score": x_value,
                        "y_external_score": y_value,
                        "reference_n_cells": x_meta["n_cells"],
                        "external_n_cells": y_meta["n_cells"],
                        "reference_library_count": x_meta["library_count"],
                        "external_library_count": y_meta["library_count"],
                        "reference_donor_count": x_meta["donor_count"],
                        "external_donor_count": y_meta["donor_count"],
                        "reference_unit_kind": x_meta["unit_kind"],
                        "external_unit_kind": y_meta["unit_kind"],
                        "external_library_is_proxy": y_meta["library_is_proxy"],
                        "reference_region_coverage": x_meta["region_coverage"],
                        "external_region_coverage": y_meta["region_coverage"],
                        "reference_condition_coverage": x_meta["condition_coverage"],
                        "external_condition_coverage": y_meta["condition_coverage"],
                        "status": status,
                    }
                )
        del frame
        gc.collect()
    table = pd.DataFrame(rows)
    table.to_csv(output_path, sep="\t", index=False, na_rep="NA", float_format="%.17g")
    print("BROAD_PROGRAM_CONSISTENCY_SOURCE_COMPLETE")
    print(f"reference_classes\t{len(reference_profiles)}")
    print(f"study_count\t{len(BROAD_CONSISTENCY_STUDIES)}")
    print(f"rows\t{len(table)}")
    print(f"finite_points\t{int(np.isfinite(table['x_reference_score']).sum() if len(table) else 0)}")
    print(f"output\t{output_path}")
    for study in BROAD_CONSISTENCY_STUDIES:
        entry = coverage[study]
        print(f"study_groups\t{study}\t{entry['groups']}\tassigned_cells\t{entry['assigned_cells']}\tclasses\t{','.join(entry['classes'])}")


def build_correlations(
    discovery: pd.DataFrame,
    profiles: dict[str, pd.DataFrame],
    program_names: list[str],
) -> None:
    profile_units = donor_group_profiles({"discovery": discovery, **profiles}, program_names)
    discovery_units = profile_units["discovery"]
    discovery_z, discovery_valid = rank_standardized(
        discovery_units[program_names].to_numpy(dtype=np.float64)
    )
    full_path = OUTPUT / "external_profile_correlations.tsv"
    same_path = OUTPUT / "external_profile_correlations_same_author_celltype.tsv"
    ranked_path = OUTPUT / "external_profile_correlations_ranked.tsv"
    with full_path.open("w", encoding="utf-8") as full, same_path.open(
        "w", encoding="utf-8"
    ):
        empty = pd.DataFrame(columns=CORRELATION_COLUMNS)
        empty.to_csv(full, sep="\t", index=False)
        empty.to_csv(same, sep="\t", index=False)
        for study in STUDIES:
            external = profile_units.get(study, pd.DataFrame(columns=profile_units["discovery"].columns))
            mapped = external[external["common_label"].ne("")].reset_index(drop=True)
            unmapped = external[external["common_label"].eq("")].reset_index(drop=True)
            if len(mapped):
                for start in range(0, len(mapped), 256):
                    chunk = mapped.iloc[start : start + 256].reset_index(drop=True)
                    external_z, external_valid = rank_standardized(
                        chunk[program_names].to_numpy(dtype=np.float64)
                    )
                    frame = paired_correlation_frame(
                        study,
                        discovery_units,
                        chunk,
                        discovery_z,
                        discovery_valid,
                        external_z,
                        external_valid,
                    )
                    frame.to_csv(full, sep="\t", index=False, header=False)
                    frame.loc[frame["same_author_celltype"].eq(True)].to_csv(
                        same, sep="\t", index=False, header=False
                    )
            if len(unmapped):
                frame = unmapped_correlation_frame(study, unmapped)
                frame.to_csv(full, sep="\t", index=False, header=False)
                frame.loc[frame["same_author_celltype"].eq(True)].to_csv(
                    same, sep="\t", index=False, header=False
                )
    write_ranked_correlations(full_path, ranked_path)
    for study, frame in profile_units.items():
        frame.to_csv(OUTPUT / f"{study}.label_profiles.tsv", sep="\t", index=False)


def main() -> None:
    parser = argparse.ArgumentParser(add_help=False)
    actions = parser.add_mutually_exclusive_group()
    actions.add_argument(
        "--postprocess-study",
        "--study",
        dest="postprocess_study",
        choices=POSTPROCESS_STUDIES,
    )
    actions.add_argument("--score-study", choices=("SEAAD",))
    actions.add_argument("--ranked-study", choices=POSTPROCESS_STUDIES)
    actions.add_argument("--merge-all", action="store_true")
    actions.add_argument("--umap-median", action="store_true")
    actions.add_argument("--median-correlations", action="store_true")
    actions.add_argument("--render-umap", action="store_true")
    actions.add_argument("--study-celltype-heatmap", action="store_true")
    actions.add_argument("--study-celltype-program-source", action="store_true")
    actions.add_argument("--annotate-study-celltype-program-source", action="store_true")
    actions.add_argument("--broad-program-consistency-source", action="store_true")
    actions.add_argument("--neuron-program-consistency", action="store_true")
    actions.add_argument("--cell-nearest-reference", action="store_true")
    args, _ = parser.parse_known_args()
    if args.cell_nearest_reference:
        OUTPUT.mkdir(parents=True, exist_ok=True)
        run_cell_nearest_reference()
        return
    if args.neuron_program_consistency:
        OUTPUT.mkdir(parents=True, exist_ok=True)
        run_neuron_program_consistency()
        return
    if args.broad_program_consistency_source:
        OUTPUT.mkdir(parents=True, exist_ok=True)
        run_broad_program_consistency_source()
        return
    if args.annotate_study_celltype_program_source:
        OUTPUT.mkdir(parents=True, exist_ok=True)
        run_annotate_study_celltype_program_source()
        return
    if args.study_celltype_program_source:
        OUTPUT.mkdir(parents=True, exist_ok=True)
        run_study_celltype_program_score_source()
        return
    if args.study_celltype_heatmap:
        OUTPUT.mkdir(parents=True, exist_ok=True)
        run_study_celltype_reference_heatmap()
        return
    if args.median_correlations:
        OUTPUT.mkdir(parents=True, exist_ok=True)
        run_library_celltype_median_correlations()
        return
    if args.render_umap:
        OUTPUT.mkdir(parents=True, exist_ok=True)
        render_library_celltype_umap()
        return
    if args.umap_median:
        OUTPUT.mkdir(parents=True, exist_ok=True)
        run_library_celltype_umap()
        return
    if args.postprocess_study:
        postprocess_one_study(args.postprocess_study)
        return
    if args.ranked_study:
        full_path = OUTPUT / f"{args.ranked_study}.profile_correlations.tsv"
        ranked_path = OUTPUT / f"{args.ranked_study}.profile_correlations_ranked.tsv"
        write_ranked_correlations(full_path, ranked_path)
        return
    if args.merge_all:
        merge_all_outputs()
        return
    if args.score_study:
        OUTPUT.mkdir(parents=True, exist_ok=True)
        sys.path.insert(0, str(CORE))
        H, ref_genes, row_number, _ = load_reference()
        old_indices, program_names, _ = load_retained_map(row_number)
        score_path, profile_path, status = score_study(
            "SEAAD", H, ref_genes, old_indices, program_names, {}
        )
        print(f"study_complete\tSEAAD\t{score_path}\t{profile_path}")
        return

    OUTPUT.mkdir(parents=True, exist_ok=True)
    sys.path.insert(0, str(CORE))
    H, ref_genes, row_number, _ = load_reference()
    old_indices, program_names, _ = load_retained_map(row_number)
    discovery = pd.read_parquet(OUTPUT / "discovery.group_profiles.parquet")
    status_path = OUTPUT / "study_status.tsv"
    existing_status = pd.read_csv(status_path, sep="\t", dtype=str).fillna("")
    pending_status = existing_status.copy()
    pending_status.loc[pending_status["study"].eq("SEAAD"), "status"] = (
        "SOURCE_ROUTING_REPAIR_IN_PROGRESS"
    )
    pending_status.loc[pending_status["study"].eq("SEAAD"), "reason"] = (
        "Re-reading approved normal/dementia cohort paths"
    )
    pending_status.to_csv(status_path, sep="\t", index=False)
    statuses = []
    group_profiles: dict[str, pd.DataFrame] = {}
    with (OUTPUT / "run.log").open("a", encoding="utf-8") as log:
        for study in STUDIES:
            score_path = OUTPUT / f"{study}.cell_scores.parquet"
            profile_path = OUTPUT / f"{study}.group_profiles.parquet"
            if study == "SEAAD":
                label_map = {
                    region: load_label_map("SEAAD", region)
                    for region in ("DLPFC", "MTG")
                }
                for cohort in ("normal", "dementia"):
                    for region in ("dlpfc", "mtg"):
                        log.write(
                            f"seaad_source\t{cohort}\t{region}\t"
                            f"{INPUT / 'SEAAD' / cohort / region}\n"
                        )
                log.flush()
                try:
                    score_path, profile_path, status = score_study(
                        study, H, ref_genes, old_indices, program_names, label_map
                    )
                    profile = pd.read_parquet(profile_path)
                    group_profiles[study] = profile
                    statuses.append(status)
                    log.write(
                        f"study_complete\t{study}\t{score_path}\t{profile_path}\n"
                    )
                    log.flush()
                except Exception as exc:
                    status = pd.DataFrame(
                        [
                            {
                                "study": study,
                                "status": "UNABLE_TO_SCORE",
                                "query_rows": "",
                                "study_rows_used_for_sd": "",
                                "score_path": str(score_path),
                                "profile_path": str(profile_path),
                                "reason": f"{type(exc).__name__}: {exc}",
                            }
                        ]
                    )
                    statuses.append(status)
                    log.write(
                        f"study_unable\t{study}\t{type(exc).__name__}: {exc}\n"
                    )
                    log.flush()
                continue
            label_map = load_label_map(study)
            profile = refresh_existing_study(
                study, score_path, profile_path, label_map, program_names
            )
            group_profiles[study] = profile
            previous = existing_status.loc[existing_status["study"].eq(study)].head(1)
            statuses.append(previous)
    status_table = pd.concat(statuses, ignore_index=True) if statuses else pd.DataFrame()
    status_table.to_csv(status_path, sep="\t", index=False)
    if all(study in group_profiles for study in STUDIES):
        build_correlations(discovery, group_profiles, program_names)


if __name__ == "__main__":
    main()
