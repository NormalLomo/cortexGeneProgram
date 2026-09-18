#!/usr/bin/env python3
"""Plot A's retained regional results. No inference or LOO is recomputed."""
from pathlib import Path
import csv, math, textwrap
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import load_workbook
ROOT=Path(__import__("os").environ["NMF_WORK_ROOT"])
DATA=ROOT/'analysis/01_existing_regional_evidence'
OUT=ROOT/'figures/human_revision/panels/fig2_donor_evidence'
S3=ROOT/'inputs/current_six_figures/Supplementary_Tables_S1-S6.xlsx'
OLD8=['P1','P3','P4','P6','P8','P9','P13','P33']

def read(n):
    with (DATA/n).open() as h:return list(csv.DictReader(h,delimiter='\t'))
def annotations():
    w=load_workbook(S3,read_only=True,data_only=True);rows=w['Table S3'].iter_rows(values_only=True);keys=next(rows)
    return {str(r[0]):dict(zip(keys,r)) for r in rows if r[0]}
def save(f,n):
    f.savefig(OUT/(n+'.pdf'),bbox_inches='tight');f.savefig(OUT/(n+'.png'),dpi=300,bbox_inches='tight',facecolor='white');plt.close(f);print('Produced',OUT/(n+'.pdf'),flush=True)
def style(ax):
    for k in ['top','right']:ax.spines[k].set_visible(False)
    ax.tick_params(length=2)

# v51-compatible replacement for panel a: donor-adjusted regional heatmap only.
# Values and q/eta annotations are read from retained tables; no inference is run.
PROGRAM_ORDER_V51 = [
    'P33','P12','P16','P23','P36','P47','P34','P49','P14','P40','P39','P53',
    'P54','P45','P32','P51','P50','P5','P25','P43','P24','P44','P52','P28',
    'P19','P11','P37','P29','P48','P21','P17','P18','P27','P30','P35','P20',
    'P46','P10','P42','P26','P41','P38','P22','P7','P31','P15','P2','P3',
    'P4','P8','P6','P13','P1','P9'
]
REGION_ORDER_V51 = ['ACC','DLPFC','AG','STG','S1','M1','V1','ITG','VLPFC','S1E','FPPFC','PoCG','SMG','SPL']
LOBE_V51 = {
    'V1':'Occipital', 'S1':'Parietal', 'S1E':'Parietal', 'PoCG':'Parietal',
    'SPL':'Parietal', 'SMG':'Parietal', 'AG':'Parietal', 'STG':'Temporal',
    'ITG':'Temporal', 'M1':'Frontal/PFC', 'VLPFC':'Frontal/PFC',
    'DLPFC':'Frontal/PFC', 'FPPFC':'Frontal/PFC', 'ACC':'Limbic'
}
LOBE_COLORS_V51 = {
    'Occipital':'#4C6EB1', 'Parietal':'#33A089', 'Temporal':'#E2A22C',
    'Frontal/PFC':'#C44E52', 'Limbic':'#8064A2'
}


def read_tsv(path):
    with Path(path).open(newline='') as handle:
        return list(csv.DictReader(handle, delimiter='\t'))


def compact_functional_label(row, program):
    name = ' '.join(str(row.get('functional_name', '')).split())
    replacements = (
        ('Regulation of ', 'Reg. '),
        ('Positive Regulation of ', 'Pos. reg. '),
        ('Negative Regulation of ', 'Neg. reg. '),
        ('Pathway', 'pathway'),
        ('Assembly', 'asm.'),
        ('Organization', 'org.'),
        ('Morphogenesis', 'morph.'),
    )
    for old, new in replacements:
        name = name.replace(old, new)
    if len(name) > 15:
        name = name[:15].rsplit(' ', 1)[0].rstrip('.,;:') + '…'
    confidence = str(row.get('confidence', '')).strip().lower()
    mark = '*' if confidence.startswith('lower') else ''
    return f'{program} {name}{mark}' if name else f'{program}{mark}'


def heatmap_only():
    profiles = read_tsv(DATA / 'HUMAN_VALIDATION_snrna_all54_donor_adjusted_region_profiles.tsv')
    robust = read_tsv(DATA / 'DONOR_ROBUST_ALL54.tsv')
    annotations = {r['new_P']: r for r in read_tsv(ROOT / 'tables' / 'TableS3_program_annotation.tsv')}
    values = {(r['program'], r['region']): float(r['standardized_effect']) for r in profiles}
    matrix = np.asarray([[values[(p, region)] for p in PROGRAM_ORDER_V51]
                         for region in REGION_ORDER_V51], dtype=float)
    stats = {r['program']: r for r in robust}
    eta = np.asarray([float(stats[p]['partial_eta_sq']) for p in PROGRAM_ORDER_V51], dtype=float)
    q = np.asarray([float(stats[p]['donor_block_permutation_F_q_all54'])
                    for p in PROGRAM_ORDER_V51], dtype=float)
    labels = [p + ('*' if str(annotations[p].get('confidence', '')).strip().lower().startswith('lower') else '')
              for p in PROGRAM_ORDER_V51]

    # Exact original panel-a canvas; no tight crop so the reducer can place it 1:1.
    fig_w_pt, fig_h_pt = 547.0, 211.0
    fig = plt.figure(figsize=(fig_w_pt / 72.0, fig_h_pt / 72.0), facecolor='white')
    left_pt, body_w_pt = 54.0, 404.0
    bottom_pt, body_h_pt = 49.0, 106.0
    body_box = [left_pt / fig_w_pt, bottom_pt / fig_h_pt,
                body_w_pt / fig_w_pt, body_h_pt / fig_h_pt]
    ax = fig.add_axes(body_box)
    im = ax.imshow(matrix, cmap='RdBu_r', vmin=-4.0, vmax=4.0,
                   aspect='auto', interpolation='nearest', origin='upper')
    ax.set_xticks(np.arange(len(PROGRAM_ORDER_V51)))
    ax.set_xticklabels(labels, rotation=90, ha='center', va='top', fontsize=6.0)
    ax.tick_params(axis='x', length=0, pad=1.0)
    ax.set_yticks(np.arange(len(REGION_ORDER_V51)))
    ax.set_yticklabels(REGION_ORDER_V51, fontsize=6.0)
    ax.tick_params(axis='y', length=0, pad=2.0)
    for spine in ax.spines.values():
        spine.set_visible(False)

    # Preserve the left lobe strip from the original hierarchy.
    lobe_rgb = np.asarray([[matplotlib.colors.to_rgb(LOBE_COLORS_V51[LOBE_V51[r]])]
                           for r in REGION_ORDER_V51], dtype=float)
    lax = fig.add_axes([(left_pt - 7.0) / fig_w_pt, bottom_pt / fig_h_pt,
                        4.0 / fig_w_pt, body_h_pt / fig_h_pt])
    lax.imshow(lobe_rgb, aspect='auto', interpolation='nearest', origin='upper')
    lax.set_xticks([]); lax.set_yticks([])
    for spine in lax.spines.values():
        spine.set_visible(False)
    fig.text((left_pt - 12.0) / fig_w_pt, (bottom_pt + body_h_pt / 2.0) / fig_h_pt,
             'lobe', rotation=90, ha='center', va='center', fontsize=5.5)

    # Existing partial-eta² summary is retained as a donor-level descriptor;
    # primary q is a separate blue/grey strip, never a confidence symbol.
    e_ax = fig.add_axes([left_pt / fig_w_pt, (bottom_pt + body_h_pt + 4.0) / fig_h_pt,
                         body_w_pt / fig_w_pt, 12.0 / fig_h_pt])
    e_ax.bar(np.arange(len(PROGRAM_ORDER_V51)), eta, width=1.0,
             color='#6E7B8B', edgecolor='none')
    e_ax.set_xlim(-0.5, len(PROGRAM_ORDER_V51) - 0.5)
    e_ax.set_ylim(0.0, max(0.25, float(np.nanmax(eta)) * 1.08))
    e_ax.set_xticks([]); e_ax.set_yticks([])
    e_ax.text(0.0, 1.02, 'partial eta^2 (donor-level)', transform=e_ax.transAxes,
              ha='left', va='bottom', fontsize=5.0, clip_on=False)
    for spine in e_ax.spines.values():
        spine.set_visible(False)

    q_rgb = np.asarray([[matplotlib.colors.to_rgb('#2166AC') if np.isfinite(x) and x < 0.05
                         else matplotlib.colors.to_rgb('#D9D9D9') for x in q]], dtype=float)
    q_ax = fig.add_axes([left_pt / fig_w_pt, (bottom_pt + body_h_pt + 19.0) / fig_h_pt,
                         body_w_pt / fig_w_pt, 4.0 / fig_h_pt])
    q_ax.imshow(q_rgb, aspect='auto', interpolation='nearest')
    q_ax.set_xticks([]); q_ax.set_yticks([])
    for spine in q_ax.spines.values():
        spine.set_visible(False)
    q_ax.text(1.002, 0.5, 'blue: primary q<0.05 (25)', transform=q_ax.transAxes,
              ha='left', va='center', fontsize=5.0)

    cax = fig.add_axes([(left_pt + body_w_pt + 11.0) / fig_w_pt, bottom_pt / fig_h_pt,
                        7.0 / fig_w_pt, body_h_pt / fig_h_pt])
    cbar = fig.colorbar(im, cax=cax, ticks=[-4, -2, 0, 2, 4])
    cbar.ax.tick_params(labelsize=5.5, length=1.5, pad=1.0)
    cbar.set_label('Standardized regional effect', fontsize=5.2, labelpad=3.0)
    fig.text(8.0 / fig_w_pt, 192.0 / fig_h_pt, 'a', fontsize=12, fontweight='bold',
             ha='left', va='center')
    fig.text(0.10, 0.965,
             'Donor-adjusted regional effects; order retained from v51 (no new clustering)',
             fontsize=5.2, ha='left', va='top')
    fig.text(0.10, 0.008,
             '* lower annotation confidence',
             fontsize=5.0, ha='left', va='bottom')

    OUT.mkdir(parents=True, exist_ok=True)
    pdf_path = OUT / 'Fig2_donor_adjusted_regional_heatmap.pdf'
    png_path = OUT / 'Fig2_donor_adjusted_regional_heatmap.png'
    fig.savefig(pdf_path, facecolor='white', bbox_inches=None, pad_inches=0)
    fig.savefig(png_path, dpi=300, facecolor='white', bbox_inches=None, pad_inches=0)
    plt.close(fig)
    print('Produced', pdf_path, png_path, flush=True)

def main():
    import sys
    if '--heatmap-only' in sys.argv:
        heatmap_only()
        return
    OUT.mkdir(parents=True,exist_ok=True)
    ann=annotations();rows=read('DONOR_ROBUST_ALL54.tsv');lookup={r['program']:r for r in rows}
    loo=read('DONOR_LOO_STABILITY.tsv');profiles=read('HUMAN_VALIDATION_snrna_all54_donor_adjusted_region_profiles.tsv');raw=read('HUMAN_VALIDATION_snrna_donor_region_pseudobulk_all54.tsv')
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'axes.titlesize':11,'pdf.fonttype':42,'ps.fonttype':42})
    # Numeric P order avoids inventing a fresh significance-ranked subset.
    ids=['P'+str(i) for i in range(1,55)]
    labs=[p+('*' if ann[p]['confidence']=='Lower confidence' else '')+
          ('†' if p in OLD8 else '') for p in ids]
    fig=plt.figure(figsize=(7.2,4.9));gs=fig.add_gridspec(1,4,width_ratios=[2.5,.55,2.5,.55],wspace=.35)
    for block in range(2):
        ax=fig.add_subplot(gs[block*2]);tab=fig.add_subplot(gs[block*2+1],sharey=ax)
        group=ids[27*block:27*(block+1)]
        for j,p in enumerate(group):
            r=lookup[p];q=float(r['donor_block_permutation_F_q_all54'])
            ax.plot([float(r['loo_partial_eta_min']),float(r['loo_partial_eta_max'])],[j,j],color='#999999',lw=.85)
            ax.scatter(float(r['partial_eta_sq']),j,s=14,color='#2166ac' if q<.05 else '#999999',zorder=3)
            for k,col in enumerate(['wild_rademacher_F_q_all54','wild_webb_F_q_all54']):
                tab.scatter(k,j,s=22,marker='s',facecolor='#b46c2b' if float(r[col])<.05 else '#e6e6e6',edgecolor='none')
        ax.set_yticks(range(27));ax.set_yticklabels([p+('*' if ann[p]['confidence']=='Lower confidence' else '')+('†' if p in OLD8 else '') for p in group],fontsize=8)
        ax.set_ylim(26.7,-.8);ax.set_xlim(0,1);ax.set_xticks([0,.5,1]);ax.tick_params(labelsize=8);ax.set_xlabel('Partial eta²',fontsize=9);style(ax)
        tab.set_xlim(-.6,1.6);tab.set_xticks([0,1]);tab.set_xticklabels(['R','W'],fontsize=8);tab.xaxis.tick_top();tab.tick_params(left=False,labelleft=False,bottom=False);tab.set_title('Sens.',fontsize=8,pad=8)
        for sp in tab.spines.values():sp.set_visible(False)
    fig.subplots_adjust(left=.075,right=.99,bottom=.22,top=.9)
    fig.text(.04,.055,'Blue: primary permutation q < 0.05. Brown: sensitivity q < 0.05 (R: Rademacher; W: Webb).',fontsize=7.5)
    fig.text(.04,.018,'Gray ranges: donor-LOO, not CI. Exact q values: Table S4. † Historical eight; * lower annotation confidence.',fontsize=7.1)
    save(fig,'Fig2_all54_effects')
    regions=sorted({r['region'] for r in profiles});pr={(r['program'],r['region']):float(r['standardized_effect']) for r in profiles}
    m=np.array([[pr[p,r] for r in regions] for p in OLD8])
    fig,ax=plt.subplots(figsize=(3.9,2.55));im=ax.imshow(m,cmap='RdBu_r',aspect='auto',vmin=-4,vmax=4)
    ax.set_xticks(range(len(regions)));ax.set_xticklabels(regions,rotation=55,ha='right',fontsize=8)
    ax.set_yticks(range(8));ax.set_yticklabels([p+('*' if ann[p]['confidence']=='Lower confidence' else '') for p in OLD8],fontsize=8)
    ax.set_title('Historical eight: adjusted profiles',loc='left',fontsize=9)
    fig.colorbar(im,ax=ax,fraction=.024,pad=.03,label='Standardized regional effect')
    fig.text(.02,.005,'Descriptive historical subset; not the final 25.',fontsize=7)
    fig.tight_layout(rect=[0,.045,1,1]);save(fig,'Fig2_original8_profiles')
    donors=sorted({r['donor'] for r in raw});cols=plt.get_cmap('tab10').colors;dc={d:cols[i%10] for i,d in enumerate(donors)}
    import sys
    selected=[(['P3','P9'],'Fig2_donor_examples')]
    if '--main-panels-only' not in sys.argv:selected.append((OLD8,'Fig2_all_original8_donor_points'))
    for chosen,stem in selected:
        nr,nc=(2,1) if len(chosen)==2 else (4,2)
        fig,axes=plt.subplots(nr,nc,figsize=(3.9 if nc==1 else 9,2.25*nr),squeeze=False)
        for ax,p in zip(axes.flat,chosen):
            for di,d in enumerate(donors):
                rr=[r for r in raw if r['donor']==d]
                x=[regions.index(r['region'])+(di-4.5)*.045 for r in rr];y=[float(r[p]) for r in rr]
                ax.scatter(x,y,s=24,c=[dc[d]],label=d,edgecolors='white',linewidths=.35)
            ax.set_xticks(range(len(regions)));ax.set_xticklabels(regions,rotation=65,ha='right',fontsize=7)
            ax.set_ylabel('Donor-region mean usage',fontsize=8.5)
            ax.set_title(p+'  '+str(ann[p]['functional_name']),loc='left',fontsize=9)
            ax.set_xlim(-.7,len(regions)-.3);style(ax)
        handles,labels=axes.flat[0].get_legend_handles_labels();fig.legend(handles,labels,loc='lower center',ncol=5,frameon=False,fontsize=7)
        fig.tight_layout(rect=[0,.085 if nr==1 else .035,1,1]);save(fig,stem)
    if '--main-panels-only' in sys.argv:return
    # Reuse the exact retained per-omission direction values, without recomputing.
    ld={(r['program'],r['omitted_donor']):float(r['centered_region_direction_agreement']) for r in loo}
    fig,ax=plt.subplots(figsize=(5.5,10.2));im=ax.imshow([[ld[p,d] for d in donors] for p in ids],aspect='auto',vmin=0,vmax=1,cmap='viridis')
    ax.set_yticks(range(54));ax.set_yticklabels(labs,fontsize=7);ax.set_xticks(range(len(donors)));ax.set_xticklabels(donors,rotation=65,ha='right',fontsize=7)
    ax.set_title('Leave-one-donor-out direction agreement\nExisting regional influence results, not independent validation',fontsize=10)
    fig.colorbar(im,ax=ax,fraction=.035,pad=.04,label='Centered regional direction agreement')
    fig.tight_layout();save(fig,'Fig2_donor_direction_sensitivity')
if __name__=='__main__':main()
