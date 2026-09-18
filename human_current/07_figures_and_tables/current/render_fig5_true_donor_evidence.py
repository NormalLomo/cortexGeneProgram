#!/usr/bin/env python3
"""Render A's true-donor spatial results without recomputing effects or LOO."""
from pathlib import Path
import csv
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import load_workbook
ROOT=Path(__import__("os").environ["NMF_WORK_ROOT"])
DATA=ROOT/'analysis/02_true_donor_spatial_support'
OUT=ROOT/'figures/human_revision/panels/fig5_true_donor_evidence'
PAIRS=['P16-P36','P16-P45','P16-P54'];DONORS=['DonorB','DonorF','DonorG','DonorH','DonorI']
COLORS={'DonorB':'#d58f00','DonorF':'#0072b2','DonorG':'#009e73','DonorH':'#cc79a7','DonorI':'#d55e00'}

def read(n):
    with (DATA/n).open() as h:return list(csv.DictReader(h,delimiter='\t'))
def save(f,n):
    f.savefig(OUT/(n+'.pdf'),bbox_inches='tight');f.savefig(OUT/(n+'.png'),bbox_inches='tight',dpi=300,facecolor='white');plt.close(f);print('Produced',OUT/(n+'.pdf'),flush=True)
def style(ax):
    for k in ['right','top']:ax.spines[k].set_visible(False)
    ax.axvline(0,color='#aaaaaa',lw=.6,zorder=0);ax.grid(axis='x',color='#eeeeee',lw=.6)
def small_legend(fig,ax):
    hh=[plt.Line2D([],[],marker='o',ls='',color=COLORS[d],label=d+' (n='+str(NS[d])+')') for d in DONORS]
    hh.append(plt.Line2D([],[],marker='D',ls='',color='black',label='All 44 sections'))
    fig.legend(handles=hh,loc='lower center',ncol=3,fontsize=8,frameon=False)

def support(per,loo,full_col,val_col,lo_col,hi_col,loo_col,name,title,xlabel):
    fig,(ax,bx)=plt.subplots(1,2,figsize=(9.4,3.9),gridspec_kw={'width_ratios':[1.35,1]})
    for j,pair in enumerate(PAIRS):
        rr=[r for r in per if r['pair']==pair];lr=[r for r in loo if r['pair']==pair]
        full=float(rr[0][full_col]);lo=float(lr[0][lo_col]);hi=float(lr[0][hi_col])
        for i,d in enumerate(DONORS):
            r=next(r for r in rr if r['donor']==d);v=float(r[val_col]);ax.scatter(v,j+(i-2)*.10,color=COLORS[d],s=28,zorder=3)
            r=next(r for r in lr if r['omitted_donor']==d);bx.scatter(float(r[loo_col]),j+(i-2)*.10,color=COLORS[d],s=22,zorder=3)
        ax.scatter(full,j,marker='D',color='black',s=26,zorder=4)
        bx.plot([lo,hi],[j,j],color='#666666',lw=1.6,zorder=1);bx.scatter(full,j,marker='D',color='black',s=24,zorder=4)
    for aa in [ax,bx]:
        aa.set_ylim(2.55,-.55);aa.set_yticks(range(3));aa.set_yticklabels(PAIRS,fontsize=10);aa.set_xlabel(xlabel,fontsize=9);style(aa)
    ax.set_title('Effects within each real donor',loc='left',fontsize=11);bx.set_title('Whole-donor omission (5 folds)',loc='left',fontsize=11)
    bx.set_yticklabels([])
    fig.suptitle(title,fontsize=12,x=.07,ha='left')
    small_legend(fig,ax)
    fig.text(.07,.155,'All-section effect is section-unweighted. Gray LOO range is influence sensitivity, not a confidence interval.',fontsize=7.8)
    fig.tight_layout(rect=[0,.20,1,.94]);save(fig,name)

def main():
    global NS
    OUT.mkdir(parents=True,exist_ok=True)
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'pdf.fonttype':42,'ps.fonttype':42})
    sb=read('same_bin_per_donor_effects.tsv');sl=read('same_bin_true_donor_loo.tsv');dp=read('P16_distance_0_500_per_donor_effects.tsv');dl=read('P16_distance_0_500_true_donor_loo.tsv');sr=read('P16_section_ring_effects.tsv')
    sb=[dict(r,pair=r['A_label']+'-'+r['B_label']) for r in sb if r['mode']=='progprog' and r['A_label']=='P16' and r['B_label'] in ['P36','P45','P54']]
    sl=[dict(r,pair=r['A_label']+'-'+r['B_label']) for r in sl if r['mode']=='progprog' and r['A_label']=='P16' and r['B_label'] in ['P36','P45','P54']]
    NS={r['donor']:int(r['n_sections_donor']) for r in dp if r['pair']==PAIRS[0]}
    support(sb,sl,'all_sections_log2_median_g','donor_log2_median_g','loo_min_log2_median_g','loo_max_log2_median_g','loo_log2_median_g','Fig5_P16_same_bin','P16 relationships within the same measurement bin','log2 median g([0,25 µm))')
    support(dp,dl,'all_sections_median_relation_0_500','donor_median_relation_0_500','loo_min_median_relation_0_500','loo_max_median_relation_0_500','loo_median_relation_0_500','Fig5_P16_distance_support','P16 relationships across 0–500 µm','Distance-weighted 0–500 µm relation')
    w=load_workbook(ROOT/'inputs/current_six_figures/Supplementary_Tables_S1-S6.xlsx',read_only=True,data_only=True);it=w['Table S3'].iter_rows(values_only=True);keys=next(it);ann={r[0]:dict(zip(keys,r)) for r in it if r[0]}
    fig,axes=plt.subplots(1,3,figsize=(10.6,3.5),sharey=True)
    for ax,pair in zip(axes,PAIRS):
        rr=[r for r in sr if r['pair']==pair];chips=sorted({r['chip'] for r in rr})
        for chip in chips:
            z=sorted((r for r in rr if r['chip']==chip),key=lambda r:float(r['ring_start_um']));don=z[0]['donor']
            # Ring centers are plotting coordinates only. The source values remain unchanged.
            xs=[(float(r['ring_start_um'])+float(r['ring_end_um']))/2 for r in z]
            ax.plot(xs,[float(r['symmetric_log2g']) for r in z],lw=.7,alpha=.35,color=COLORS[don])
        ax.axhline(0,color='#999999',lw=.7);ax.set_xlim(0,500);ax.set_xticks([0,100,250,500]);ax.set_xlabel('Distance ring (µm)')
        ax.set_title(pair+'\n'+str(ann[pair.split('-')[1]]['functional_name']),fontsize=9)
        for k in ['right','top']:ax.spines[k].set_visible(False)
    axes[0].set_ylabel('Section-specific symmetric log2 g(r)')
    handles=[plt.Line2D([],[],color=COLORS[d],label=d+' (n='+str(NS[d])+')') for d in DONORS]
    fig.legend(handles=handles,loc='lower center',ncol=5,frameon=False,fontsize=8)
    fig.suptitle('All 44 sections: spatial relationships across ten distance rings',fontsize=12,x=.07,ha='left')
    fig.tight_layout(rect=[0,.10,1,.90]);save(fig,'Fig5_P16_section_distance_profiles')
    fig,axes=plt.subplots(2,3,figsize=(10.4,6.4),sharex=True)
    for j,pair in enumerate(PAIRS):
        rr=[r for r in sr if r['pair']==pair and float(r['ring_start_um'])==0 and float(r['ring_end_um'])==25]
        for di,d in enumerate(DONORS):
            z=sorted([r for r in rr if r['donor']==d],key=lambda r:r['chip'])
            for k,r in enumerate(z):
                x=di+(k-(len(z)-1)/2)*.015
                axes[0,j].scatter(x,float(r['symmetric_log2g']),s=21,color=COLORS[d],alpha=.8)
                # Exactly one section record is used, not its ten repeated ring records.
                axes[1,j].scatter(x,float(r['section_relation_0_500']),s=21,color=COLORS[d],alpha=.8)
        axes[0,j].set_title(pair,fontsize=11)
        for row in range(2):
            ax=axes[row,j];ax.axhline(0,color='#999999',lw=.7);ax.set_xticks(range(5));ax.set_xticklabels([d.replace('Donor','')+' (n='+str(NS[d])+')' for d in DONORS],rotation=30,ha='right',fontsize=8)
            for side in ['right','top']:ax.spines[side].set_visible(False)
    axes[0,0].set_ylabel('Same-bin symmetric log2 g');axes[1,0].set_ylabel('0–500 µm section relation')
    fig.suptitle('Section-level observations grouped by the five real donors',fontsize=12)
    fig.text(.04,.01,'Each dot is one section. Unequal donor coverage: 1 / 10 / 13 / 16 / 4 sections; no bin-level pseudoreplication.',fontsize=8)
    fig.tight_layout(rect=[0,.045,1,.95]);save(fig,'Fig5_P16_section_points')
if __name__=='__main__':main()
